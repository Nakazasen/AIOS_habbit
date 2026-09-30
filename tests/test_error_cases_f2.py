"""F2 tests: importer for the 29-column History KDTPS sheet.

Uses a small synthetic xlsx fixture (not the full 15.737-row file).
Covers: column mapping (S=part code not handler, line=E), dedup on
(year, NO, machine, line), skipped rows (missing year/NO/machine),
green-cell code path, and "vat lai file cu thi bo qua" (re-import skip).
"""
import os
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from aios_habit.error_cases import (
    HISTORY_SHEET_NAME,
    column_map,
    fill_occurred_at,
    import_history,
    parse_date_cell,
    store,
)

GREEN_FILL = PatternFill(
    start_color="FF92D050", end_color="FF92D050", fill_type="solid"
)  # RGB(146, 208, 80)


def make_row(over=None):
    """29-col history row; kwargs are 0-based column indexes."""
    cells = [None] * 29
    base = {
        0: 2024, 1: 1, 2: "2024-01-05", 3: "Virgo", 4: "C33",
        6: "SER123", 7: "ERROR", 8: "desc", 9: 1,
        11: "ctx jp", 12: "ctx vn", 13: "JP inv", 14: "VN inv",
        15: "ME", 16: "2024-02-01", 17: "part name", 18: "3VC0Y01020",
        19: "LOT1", 20: "NCC1", 21: "×", 22: "no", 23: "2024-01-06",
        24: "ー", 25: "有", 26: "cause", 27: "kaizen", 28: "http://x",
    }
    base.update(over or {})
    for i, v in base.items():
        cells[i] = v
    return cells


@pytest.fixture()
def fixture_xlsx(tmp_path):
    path = tmp_path / "Loi KDTPS.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = HISTORY_SHEET_NAME
    for c in range(29):  # header row 4
        ws.cell(row=4, column=c + 1, value=f"COL{c}")
    rows = [
        make_row(),                                            # row 5: normal
        make_row({3: "6th A4", 4: "C25"}),                    # row 6: same (year,NO), other machine/line
        make_row({0: None, 1: 5}),                            # row 7: missing year -> skip
        make_row({0: 2024, 1: None}),                         # row 8: missing NO -> skip
        make_row({0: 2024, 1: 2, 3: None}),                   # row 9: missing machine -> skip
        make_row({0: 2024, 1: 3, 14: "green me"}),            # row 10: green O cell
        make_row({0: 2025, 1: 1, 3: "Polaris", 4: "C21"}),    # row 11: normal
    ]
    for r, cells in enumerate(rows, start=5):
        for c, v in enumerate(cells):
            ws.cell(row=r, column=c + 1, value=v)
    ws.cell(row=10, column=15).fill = GREEN_FILL  # O = investigation VN
    wb.save(path)
    wb.close()
    return path


@pytest.fixture()
def conn():
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    store.init_db(c)
    yield c
    c.close()


def test_import_counts(conn, fixture_xlsx):
    r = import_history(conn, fixture_xlsx)
    assert r["status"] == "imported"
    assert r["rows_read"] == 7
    assert r["rows_imported"] == 4
    assert r["rows_skipped"] == 3
    assert store.count_cases(conn) == 4
    stats = store.batch_stats(conn, r["batch_id"])
    assert stats["sheet_name"] == HISTORY_SHEET_NAME
    assert stats["header_row"] == 4
    assert stats["source_file"] == "Loi KDTPS.xlsx"


def test_mapping_columns(conn, fixture_xlsx):
    import_history(conn, fixture_xlsx)
    row = conn.execute(
        "SELECT * FROM error_cases WHERE no_dvd='2024/1' AND machine_type='Virgo'"
    ).fetchone()
    assert row is not None
    assert row["machine_type"] == "Virgo"   # D, not C
    assert row["line"] == "C33"             # E is the true line
    assert row["error_code_h"] == "ERROR"
    assert row["department"] == "ME"
    assert row["investigation"] == "VN inv"  # VN preferred over JP
    assert row["handler"] is None            # S is a part code, not a person
    assert row["is_completed"] == ""         # V=LKATQT, not a completion flag
    assert row["needs_jp_support"] == ""
    import json
    raw = json.loads(row["raw_json"])
    assert raw["S"] == "3VC0Y01020"


def test_dedup_same_year_no_other_machine_line(conn, fixture_xlsx):
    """Same (year, NO) with different machine/line -> two separate rows."""
    import_history(conn, fixture_xlsx)
    rows = conn.execute(
        "SELECT machine_type, line FROM error_cases WHERE no_dvd='2024/1'"
    ).fetchall()
    assert {(r["machine_type"], r["line"]) for r in rows} == {
        ("Virgo", "C33"), ("6th A4", "C25")
    }


def test_green_cell_code_path(conn, fixture_xlsx):
    import_history(conn, fixture_xlsx)
    row = conn.execute(
        "SELECT skip_cells FROM error_cases WHERE no_dvd='2024/3'"
    ).fetchone()
    import json
    assert "O" in json.loads(row["skip_cells"])


def test_reimport_same_file_skipped(conn, fixture_xlsx):
    """'Vat lai file cu thi bo qua': unchanged file -> skipped, no writes."""
    first = import_history(conn, fixture_xlsx)
    assert first["status"] == "imported"
    before = store.count_cases(conn)
    batches_before = conn.execute(
        "SELECT COUNT(*) n FROM import_batches").fetchone()["n"]
    second = import_history(conn, fixture_xlsx)
    assert second["status"] == "skipped"
    assert second["batch_id"] == first["batch_id"]
    assert store.count_cases(conn) == before
    assert conn.execute(
        "SELECT COUNT(*) n FROM import_batches").fetchone()["n"] == batches_before


def test_force_reimport_refreshes(conn, fixture_xlsx):
    first = import_history(conn, fixture_xlsx)
    r = import_history(conn, fixture_xlsx, force=True)
    assert r["status"] == "imported"
    assert r["batch_id"] == first["batch_id"]  # same batch row refreshed
    assert r["rows_imported"] == 4
    assert store.count_cases(conn) == 4  # upserts, no duplicates


# ---------------------------------------------------------------------------
# date-map (ticket): occurred_at from column C (production date)
# ---------------------------------------------------------------------------

def test_parse_date_cell_formats_and_placeholders():
    from datetime import datetime
    assert parse_date_cell(datetime(2024, 1, 5)) == "2024-01-05"
    assert parse_date_cell("2024-01-05") == "2024-01-05"
    assert parse_date_cell("2024/1/5") == "2024-01-05"
    assert parse_date_cell("2024.01.05") == "2024-01-05"
    assert parse_date_cell("2024年1月5日") == "2024-01-05"
    assert parse_date_cell("2024/1/5日") == "2024-01-05"
    for empty in (None, "", "  ", "ー", "-", "abc", "2/15日"):
        assert parse_date_cell(empty) is None, empty


def test_occurred_at_extracted_from_column_c(conn, fixture_xlsx):
    r = import_history(conn, fixture_xlsx)
    assert r["rows_with_occurred_at"] == 4  # every importable fixture row has C
    row = conn.execute(
        "SELECT occurred_at, created_at FROM error_cases"
        " WHERE no_dvd='2024/1' AND machine_type='Virgo'"
    ).fetchone()
    assert row["occurred_at"] == "2024-01-05"          # C, not X/Y
    assert row["created_at"] != row["occurred_at"]     # import time untouched


def test_occurred_at_null_when_source_date_missing(conn, tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "Loi KDTPS.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = HISTORY_SHEET_NAME
    for c in range(29):
        ws.cell(row=4, column=c + 1, value=f"COL{c}")
    cells = make_row()
    cells[2] = None  # C missing in the source
    for c, v in enumerate(cells):
        ws.cell(row=5, column=c + 1, value=v)
    wb.save(path)
    wb.close()
    r = import_history(conn, path)
    assert r["rows_with_occurred_at"] == 0
    row = conn.execute("SELECT occurred_at FROM error_cases").fetchone()
    assert row["occurred_at"] is None


def test_fill_occurred_at_plan_apply_idempotent(conn, fixture_xlsx):
    import_history(conn, fixture_xlsx)
    conn.execute("UPDATE error_cases SET occurred_at = NULL")  # DB built pre-migration
    conn.commit()

    plan = fill_occurred_at(conn, fixture_xlsx)  # dry-run by default
    assert plan["status"] == "planned"
    assert plan["would_update"] == 4 and plan["updated"] == 0
    assert conn.execute(
        "SELECT COUNT(*) n FROM error_cases WHERE occurred_at IS NULL").fetchone()["n"] == 4

    applied = fill_occurred_at(conn, fixture_xlsx, apply=True)
    assert applied["status"] == "applied"
    assert applied["updated"] == 4
    assert conn.execute(
        "SELECT COUNT(*) n FROM error_cases WHERE occurred_at IS NULL").fetchone()["n"] == 0
    row = conn.execute(
        "SELECT occurred_at FROM error_cases WHERE no_dvd='2024/3'").fetchone()
    assert row["occurred_at"] == "2024-01-05"

    again = fill_occurred_at(conn, fixture_xlsx, apply=True)
    assert again["would_update"] == 0 and again["no_change"] == 4


def test_fill_occurred_at_never_touches_mismatched_rows(conn, fixture_xlsx):
    import_history(conn, fixture_xlsx)
    conn.execute("UPDATE error_cases SET occurred_at = NULL")
    conn.execute(
        "UPDATE error_cases SET line = 'CHANGED'"
        " WHERE no_dvd='2024/1' AND machine_type='Virgo'"
    )
    conn.commit()
    r = fill_occurred_at(conn, fixture_xlsx, apply=True)
    assert r["mismatch"] == 1
    assert r["updated"] == 3
    row = conn.execute(
        "SELECT occurred_at FROM error_cases"
        " WHERE no_dvd='2024/1' AND machine_type='Virgo'").fetchone()
    assert row["occurred_at"] is None  # mismatched row untouched


def test_fill_occurred_at_unknown_file_reports_no_batch(conn, fixture_xlsx, tmp_path):
    import_history(conn, fixture_xlsx)
    other = tmp_path / "other_name.xlsx"
    other.write_bytes(fixture_xlsx.read_bytes())
    r = fill_occurred_at(conn, other)
    assert r["status"] == "no_batch"
    assert r["batch_id"] is None


REAL_FILE = (
    Path(os.environ.get("AIOS_DATA_DIR", str(Path.home() / "workspace/aios_data")))
    / "dieu_tra_loi" / "Điều chỉnh" / "Lịch sử lỗi" / "Loi KDTPS.xlsx"
)


@pytest.mark.skipif(not REAL_FILE.exists(), reason="real source file not on this machine")
def test_real_file_first_rows_parse():
    """Smoke check: column assumptions hold on the real file (read-only)."""
    import openpyxl
    wb = openpyxl.load_workbook(REAL_FILE, read_only=True, data_only=True)
    try:
        ws = wb[HISTORY_SHEET_NAME]
        rows = list(ws.iter_rows(min_row=5, max_row=7, max_col=29, values_only=True))
        assert len(rows) == 3
        for values in rows:
            f = column_map.normalize_history_row(list(values))
            year, no = str(values[0]).strip(), str(values[1]).strip()
            assert f["no_dvd"] == f"{year}/{no}"
            assert f["machine_type"]  # D
            assert f["line"]          # E
    finally:
        wb.close()
