"""Tests for the investigation-report chat action (agent báo cáo điều tra).

Seeded mini-DB (deterministic): four same-code cases so the report must
contain the real countermeasure from history, trigger the recurrence
alert, span >= 3 weeks for the trend chart — and must NOT contain any
invented error code or invented countermeasure.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from docx import Document

from aios_habit import chat_action
from aios_habit.chat_action import (
    BLOCK_CHART,
    ChatActionRequest,
    match_action,
    render_outcome,
)
from aios_habit.chat_action_bao_cao_dieu_tra import ACTION_NAME, _handler
from aios_habit.error_cases import connect, init_db
from aios_habit.feature_flags import (
    FEATURE_INVESTIGATION_REPORT,
    override_feature_flags,
    reset_feature_flags,
)

TARGET_TICKET = "FORM-20261001-0001"
CODE = "C7620"
REAL_REMEDIES = ("Cắm lại cáp FFC vào mainboard", "Thay cáp FFC mới")


@pytest.fixture(autouse=True)
def _clean_state():
    reset_feature_flags()
    chat_action.reset_actions()
    yield
    reset_feature_flags()
    chat_action.reset_actions()


def _seed_db(tmp_path: Path) -> Path:
    db_path = tmp_path / "error_cases.db"
    conn = connect(str(db_path))
    init_db(conn)
    conn.execute(
        "INSERT INTO import_batches "
        "(source_file, file_sha256, sheet_name, imported_at) VALUES (?, ?, ?, ?)",
        ("Loi KDTPS.xlsx", "deadbeef", "History KDTPS", "2026-10-01 09:00:00"),
    )
    cases = [
        # ticket, occurred, phenomenon, cause, countermeasure
        (TARGET_TICKET, "2026-10-01",
         "LCD không hiển thị khi bật nguồn", "Lỏng cáp FFC nối LCD với mainboard",
         "Thay cáp FFC nối LCD"),
        ("2023/5", "2026-09-30",
         "Màn hình LCD tối đen sau khi khởi động", "Cáp FFC tuột khỏi đầu nối",
         REAL_REMEDIES[0]),
        ("2023/6", "2026-09-15",
         "LCD không lên hình", "Cáp FFC bị đứt ngầm",
         REAL_REMEDIES[1]),
        ("2023/7", "2026-09-01",
         "Màn hình chập chờn rồi tắt", "Đầu nối LCD oxy hóa",
         "Vệ sinh đầu nối LCD"),
    ]
    for idx, (ticket, occurred, phen, cause, remedy) in enumerate(cases, start=1):
        raw = {"I": phen, "AA": cause, "AB": remedy}
        conn.execute(
            "INSERT INTO error_cases "
            "(batch_id, source_row, no_dvd, machine_type, line, error_code_c, "
            " investigation, occurred_at, raw_json, phenomenon, cause, "
            " countermeasure, process_stage, department) "
            "VALUES (1, ?, ?, 'M4080', 'Line 1', ?, ?, ?, ?, ?, ?, ?, 'In ấn', 'Kỹ thuật')",
            (100 + idx, ticket, CODE,
             f"Điều tra {ticket}", occurred, json.dumps(raw, ensure_ascii=False),
             phen, cause, remedy),
        )
    conn.commit()
    conn.close()
    return db_path


def _request(question: str, db_path: Path, out_dir: Path) -> ChatActionRequest:
    return ChatActionRequest(
        question=question,
        context={
            "error_cases_db": str(db_path),
            "report_out_dir": str(out_dir),
        },
    )


def _on():
    return override_feature_flags(**{FEATURE_INVESTIGATION_REPORT: True})


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _outcome_text(outcome) -> str:
    return render_outcome(outcome)


# ---------------------------------------------------------------------------
# Flag gate
# ---------------------------------------------------------------------------


def test_flag_off_returns_none(tmp_path):
    db = _seed_db(tmp_path)
    req = _request(f"lập báo cáo điều tra cho ca {TARGET_TICKET}", db, tmp_path / "out")
    assert _handler(req) is None


def test_missing_db_returns_none(tmp_path):
    with _on():
        req = _request(
            f"lập báo cáo điều tra cho ca {TARGET_TICKET}",
            tmp_path / "khong-co.db",
            tmp_path / "out",
        )
    assert _handler(req) is None


# ---------------------------------------------------------------------------
# Happy path
# ---------------------------------------------------------------------------


def test_report_by_ticket_number(tmp_path):
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    with _on():
        outcome = _handler(_request(
            f"lập báo cáo điều tra cho ca {TARGET_TICKET}", db, out_dir))
    assert outcome is not None
    assert outcome.action == ACTION_NAME
    text = _outcome_text(outcome)

    docx_files = sorted(out_dir.glob("*.docx"))
    assert len(docx_files) == 1
    docx_path = docx_files[0]
    assert str(docx_path) in text

    body = _docx_text(docx_path)
    # All report sections present.
    for section in (
        "1. Thông tin ca",
        "2. Từ điển mã lỗi",
        "3. Ca tương tự trong lịch sử",
        "4. Đối sách đã dùng ở ca tương tự",
        "5. Cây điều tra 4M",
        "6. Phân loại tự động",
        "7. Tái phát & xu hướng",
    ):
        assert section in body, section
    # Real facts from the seeded DB.
    assert TARGET_TICKET in body
    assert CODE in body
    for remedy in REAL_REMEDIES:
        assert remedy in body
    # Recurrence alert fires (same code within 168h around the case).
    assert "CẢNH BÁO TÁI PHÁT" in body
    # Missing-data honesty: no glossary was seeded.
    assert "Chưa có dữ liệu" in body
    # No invented error codes anywhere in the document.
    codes_in_doc = set(re.findall(r"\b[CFJ]\d{3,4}\b", body))
    assert codes_in_doc <= {CODE}, codes_in_doc


def test_report_by_error_code_picks_latest_case(tmp_path):
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    with _on():
        outcome = _handler(_request(
            "lập báo cáo điều tra cho ca C7620", db, out_dir))
    assert outcome is not None
    text = _outcome_text(outcome)
    assert TARGET_TICKET in text  # latest occurred_at wins
    docx_files = sorted(out_dir.glob("*.docx"))
    assert len(docx_files) == 1
    assert TARGET_TICKET in docx_files[0].name


def test_trend_chart_file_and_block(tmp_path):
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    with _on():
        outcome = _handler(_request(
            f"lập báo cáo điều tra cho ca {TARGET_TICKET}", db, out_dir))
    assert outcome is not None
    # 3 distinct weeks of same-code cases -> chart PNG saved + inline block.
    pngs = sorted(out_dir.glob("*.png"))
    assert len(pngs) == 1
    assert pngs[0].stat().st_size > 0
    assert any(b.kind == BLOCK_CHART for b in outcome.blocks)


# ---------------------------------------------------------------------------
# Unknown / unspecified case: guidance, never invention
# ---------------------------------------------------------------------------


def test_unknown_ticket_guides_without_creating_file(tmp_path):
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    with _on():
        outcome = _handler(_request(
            "lập báo cáo điều tra cho ca FORM-20269999-9999", db, out_dir))
    assert outcome is not None
    assert "Không tìm thấy ca" in _outcome_text(outcome)
    assert list(out_dir.glob("*.docx")) == []


def test_unspecified_case_asks_for_identifier(tmp_path):
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    with _on():
        outcome = _handler(_request("lập báo cáo điều tra", db, out_dir))
    assert outcome is not None
    text = _outcome_text(outcome)
    assert "số phiếu" in text
    assert list(out_dir.glob("*.docx")) == []


# ---------------------------------------------------------------------------
# Routing priority: the command must not be swallowed by other actions
# ---------------------------------------------------------------------------


def test_match_priority_over_bare_code_lookup():
    chat_action.load_builtin_actions()
    action = match_action(ChatActionRequest(
        question="lập báo cáo điều tra cho ca C7620"))
    assert action is not None
    assert action.name == ACTION_NAME
    # And the plain investigation-plan command still routes to Buoc 3.
    plan = match_action(ChatActionRequest(
        question="gợi ý hướng điều tra cho lỗi kẹt giấy"))
    assert plan is not None
    assert plan.name == "goi_y_huong_dieu_tra"


def test_readonly_db_not_modified(tmp_path):
    db = _seed_db(tmp_path)
    with open(db, "rb") as fh:
        before = fh.read()
    with _on():
        _handler(_request(f"lập báo cáo điều tra cho ca {TARGET_TICKET}",
                            db, tmp_path / "out"))
    with open(db, "rb") as fh:
        after = fh.read()
    assert before == after


def test_existing_report_file_does_not_block_new_one(tmp_path):
    """Re-running the command simply writes a fresh timestamped report."""
    db = _seed_db(tmp_path)
    out_dir = tmp_path / "out"
    req = _request(f"lập báo cáo điều tra cho ca {TARGET_TICKET}", db, out_dir)
    with _on():
        first = _handler(req)
        second = _handler(req)
    assert first is not None and second is not None
    assert len(list(out_dir.glob("*.docx"))) >= 1
