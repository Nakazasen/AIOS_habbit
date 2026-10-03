"""Tests cho UI phuong an A ve agent-report (Task 2):

- Nut "Tai ve" nhan biet dinh dang: .docx/.pptx -> nhi phan + MIME dung;
  .md -> van ban + xem truoc.
- Nut "Hoan tac" dung file sao luu backend (checkpoint_path): sua -> khoi
  phuc ban goc; tao moi (khong backup) -> xoa file vua tao; tu choi duong
  dan khong an toan.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from aios_habit.agent_report_artifact import (
    ARTIFACT_MIME,
    artifact_download_payload,
    is_binary_artifact_suffix,
    restore_report_backup,
)


def _write(path: Path, data) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, bytes):
        path.write_bytes(data)
    else:
        path.write_text(data, encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# Download payload
# ---------------------------------------------------------------------------


def test_is_binary_artifact_suffix():
    assert is_binary_artifact_suffix(".docx") is True
    assert is_binary_artifact_suffix(".pptx") is True
    assert is_binary_artifact_suffix(".DOCX") is True
    assert is_binary_artifact_suffix(".md") is False
    assert is_binary_artifact_suffix(".txt") is False


def test_payload_md_is_text(tmp_path):
    target = _write(tmp_path / "bao-cao.md", "# Tieu de\n\nNoi dung.")
    payload = artifact_download_payload(target)
    assert payload["ok"] is True
    assert isinstance(payload["data"], str)
    assert "Noi dung" in payload["data"]
    assert payload["mime"] == "text/markdown"
    assert payload["file_name"] == "bao-cao.md"
    assert payload["is_binary"] is False


def test_payload_docx_is_binary_with_word_mime(tmp_path):
    blob = b"PK\x03\x04" + b"\x00" * 100  # zip header, content irrelevant here
    target = _write(tmp_path / "bao-cao.docx", blob)
    payload = artifact_download_payload(target)
    assert payload["ok"] is True
    assert isinstance(payload["data"], bytes)
    assert payload["data"] == blob
    assert payload["mime"] == ARTIFACT_MIME[".docx"]
    assert "wordprocessingml" in payload["mime"]
    assert payload["file_name"] == "bao-cao.docx"
    assert payload["is_binary"] is True


def test_payload_pptx_is_binary_with_powerpoint_mime(tmp_path):
    blob = b"PK\x03\x04" + b"\x11" * 100
    target = _write(tmp_path / "trinh-chieu.pptx", blob)
    payload = artifact_download_payload(target)
    assert payload["ok"] is True
    assert isinstance(payload["data"], bytes)
    assert payload["mime"] == ARTIFACT_MIME[".pptx"]
    assert "presentationml" in payload["mime"]
    assert payload["is_binary"] is True


def test_payload_missing_file_fails_vietnamese(tmp_path):
    payload = artifact_download_payload(tmp_path / "khong-co.md")
    assert payload["ok"] is False
    assert "Chưa tìm thấy file" in payload["error_vi"]


# ---------------------------------------------------------------------------
# Undo from the backend backup
# ---------------------------------------------------------------------------


def test_restore_edit_recovers_original(tmp_path):
    target = _write(tmp_path / "bao-cao.md", "BAN GOC")
    backup = _write(tmp_path / "bao-cao.md.bak-1", "BAN GOC")
    target.write_text("DA SUA", encoding="utf-8")
    ok, message = restore_report_backup(target, backup)
    assert ok is True
    assert target.read_text(encoding="utf-8") == "BAN GOC"
    assert "hoàn tác" in message


def test_restore_create_without_backup_deletes_file(tmp_path):
    target = _write(tmp_path / "moi.md", "NOI DUNG MOI")
    ok, message = restore_report_backup(target, "")
    assert ok is True
    assert not target.exists()
    assert "xóa file vừa tạo" in message


def test_restore_missing_checkpoint_fails(tmp_path):
    target = _write(tmp_path / "bao-cao.md", "DA SUA")
    ok, message = restore_report_backup(target, tmp_path / "khong-co.bak")
    assert ok is False
    assert "Chưa tìm thấy file sao lưu" in message
    assert target.read_text(encoding="utf-8") == "DA SUA"


def test_restore_checkpoint_in_other_dir_refused(tmp_path):
    target = _write(tmp_path / "a" / "bao-cao.md", "DA SUA")
    backup = _write(tmp_path / "b" / "bao-cao.md.bak", "BAN GOC")
    ok, message = restore_report_backup(target, backup)
    assert ok is False
    assert "bị từ chối" in message
    assert target.read_text(encoding="utf-8") == "DA SUA"


def test_restore_traversal_refused(tmp_path):
    evil = tmp_path / ".." / "bao-cao.md"
    backup = _write(tmp_path / "bao-cao.md.bak", "BAN GOC")
    ok, message = restore_report_backup(str(evil), backup)
    assert ok is False
    assert "bị từ chối" in message


def test_restore_unsupported_suffix_refused(tmp_path):
    target = _write(tmp_path / "run.exe", "x")
    backup = _write(tmp_path / "run.exe.bak", "x")
    ok, message = restore_report_backup(target, backup)
    assert ok is False
    assert "bị từ chối" in message


def test_restore_empty_result_path_fails():
    ok, message = restore_report_backup("", "")
    assert ok is False
    assert "Thiếu đường dẫn" in message


# ---------------------------------------------------------------------------
# End-to-end with the real backend: edit -> backup -> undo restores bytes
# ---------------------------------------------------------------------------


def test_e2e_docx_edit_then_undo_restores_binary(tmp_path, monkeypatch):
    """Tao docx that bang backend, sua, hoan tac -> SHA goc duoc khoi phuc."""
    pytest.importorskip("docx")
    monkeypatch.setenv("AIOS_LOCAL_CASES_DIR", str(tmp_path))
    monkeypatch.setenv("AIOS_DOC_ROOT", str(tmp_path))
    from docx import Document

    from aios_habit.chat_action_agent_report import run_report_edit

    created = run_report_edit(
        "e2e.docx",
        [{"op": "append_paragraph", "text": "Doan goc"}],
        create=True,
    )
    assert created["ok"] is True
    original_sha = created["sha256"]

    edited = run_report_edit(
        "e2e.docx",
        [{"op": "append_paragraph", "text": "Doan them"}],
        create=False,
    )
    assert edited["ok"] is True
    assert edited["backup"], "backend phai tao backup truoc khi sua"
    assert len(Document(str(edited["path"])).paragraphs) == 2

    # Nut "Tai ve" phai nhan ra file nhi phan.
    payload = artifact_download_payload(edited["path"])
    assert payload["ok"] is True and payload["is_binary"] is True

    # Nut "Hoan tac" khoi phuc dung ban goc.
    ok, _ = restore_report_backup(edited["path"], edited["backup"])
    assert ok is True
    assert len(Document(str(edited["path"])).paragraphs) == 1

    import hashlib

    assert hashlib.sha256(Path(edited["path"]).read_bytes()).hexdigest() == original_sha
