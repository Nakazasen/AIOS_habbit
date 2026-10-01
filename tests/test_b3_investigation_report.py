"""B3 tests: investigation plan from the company KTD report template.

DỮ LIỆU THẬT — REAL DATA.
The three phenomena below are copied from column I
("不具合現象 / Hiện trạng lỗi") of the real "History KDTPS" sheet in
Loi KDTPS.xlsx (the 15,707 real error cases), rows 20 / 34 / 16.
They are content-verbatim with ONE deliberate whitespace difference:
the real cell for REAL_PHENOMENON_CODED holds a line break between the
panel text and the code ("LCD画面に2340表示\\nC2340：…"); the constant
below renders that break as a space so the string stays single-line
(see docs/phieu-viec/ket-qua/b3.md section 3.2). Everything else is
byte-identical to the source cells.

- REAL_PHENOMENON_CODED: C CALL case, panel shows C2340
  (定着圧解除モータエラー/timeout).
- REAL_PHENOMENON_TORN_CASE: appearance case, torn outer case.
- REAL_PHENOMENON_STOPPER: function case, stopper paper not moving smoothly.

Per the B3 ticket the generator must produce, for each of them: all four
4M branches, a Why-Why chain at least 3 levels deep, and specific
confirmation items (never vague filler like "kiểm tra lại máy").
"""

import pytest

from aios_habit.error_cases.investigation_tree import (
    BRANCHES,
    REPORT_FIELDS,
    REPORT_TITLE,
    build_tree,
    export_docx,
    export_report,
    render_report,
)

# --- verbatim real phenomena (History KDTPS, column I) -----------------------
REAL_PHENOMENON_CODED = (
    "LCD画面に2340表示 C2340：定着圧解除モータエラー（タイムアウト）"
)
REAL_PHENOMENON_TORN_CASE = "OUTER CASE-V LP-S280DN EHCが破れ"
REAL_PHENOMENON_STOPPER = "Stopper Paperがスムーズに動かない"

REAL_PHENOMENA = (
    REAL_PHENOMENON_CODED,
    REAL_PHENOMENON_TORN_CASE,
    REAL_PHENOMENON_STOPPER,
)

# Vague filler the ticket explicitly forbids ("không chung chung kiểu
# 'kiểm tra lại máy'").
VAGUE_PHRASES = (
    "kiểm tra lại máy",
    "kiểm tra chung",
    "xem lại máy",
    "kiểm tra tổng quát",
    "rà soát lại",
)


@pytest.mark.parametrize("phenomenon", REAL_PHENOMENA)
def test_b3_real_phenomenon_tree_has_all_4m(phenomenon):
    tree = build_tree(phenomenon)
    assert tree.branch_order == list(BRANCHES)
    for branch in BRANCHES:
        items = tree.branches[branch]
        assert len(items) >= 4, branch


@pytest.mark.parametrize("phenomenon", REAL_PHENOMENA)
def test_b3_real_phenomenon_why_why_at_least_3_levels(phenomenon):
    tree = build_tree(phenomenon)
    assert len(tree.why_chain) >= 3
    assert [n.level for n in tree.why_chain] == list(
        range(1, len(tree.why_chain) + 1))
    assert all(n.hint for n in tree.why_chain)
    assert phenomenon in tree.why_chain[0].question


@pytest.mark.parametrize("phenomenon", REAL_PHENOMENA)
def test_b3_real_phenomenon_items_are_specific(phenomenon):
    tree = build_tree(phenomenon)
    for item in tree.all_items():
        assert item.question and item.question.strip()
        # Every item names concrete data/artifacts to collect.
        assert item.data_to_collect and item.data_to_collect.strip()
        lowered_q = item.question.lower()
        lowered_d = item.data_to_collect.lower()
        for vague in VAGUE_PHRASES:
            assert vague not in lowered_q, item.question
            assert vague not in lowered_d, item.data_to_collect


def test_b3_coded_phenomenon_extracts_real_code():
    # The C2340 token is a real displayed code (verbatim from the case).
    from aios_habit.error_cases.column_map import extract_code_from_text
    assert extract_code_from_text(REAL_PHENOMENON_CODED) == "C2340"
    tree = build_tree(REAL_PHENOMENON_CODED, code="C2340")
    assert "C2340" in render_report(tree)


def test_b3_jam_phenomenon_gets_targeted_items():
    tree = build_tree("パネル画面にJAM 9110表示")
    machine_qs = [it.question for it in tree.branches["Machine"]]
    assert any("Vị trí kẹt" in q for q in machine_qs)


# ---------------------------------------------------------------------------
# Company report template
# ---------------------------------------------------------------------------

def test_b3_report_uses_verbatim_company_fields():
    tree = build_tree(REAL_PHENOMENON_CODED, code="C2340")
    md = render_report(tree)
    assert REPORT_TITLE in md  # Báo cáo điều tra lỗi/調査報告書
    labels = [label for _, label in REPORT_FIELDS]
    for label in labels:
        if label.startswith("Investigation content"):
            continue
        # Two-line KTD labels render their "\n" as <br> inside the table.
        assert label.replace("\n", "<br>") in md, label
    # The plan lands in the Investigation content and results section.
    assert "Investigation content and results" in md
    assert "Cây điều tra 4M" in md
    assert REAL_PHENOMENON_CODED in md  # Contents of defect seeded


def test_b3_labels_byte_match_real_ktd_cells():
    # Byte-exact labels vs the real KTD "Bao cao dieu tra" cells, reprs
    # copied from docs/phieu-viec/ket-qua/b3.md section 3.1 (verified on
    # all 215 real KTD files): U+3000 ideographic spaces, real "\n" in the
    # two-line labels, space before ")" and trailing ":" restored.
    real_cells = {
        "machine_no": "Machine\u3000No.／仕上げ-マシンNo.",
        "occurrence_date": "Occurrence\u3000Date／発生日",
        "status_at_line": "Status of occurrence at Line\nラインでの発生状況",
        "reappear_rate": "Reappear rate(%)\n(Describe the reappear environment )",
        "investigation": "Investigation content and results\n調査内容と結果:",
    }
    labels = dict(REPORT_FIELDS)
    assert set(real_cells) <= set(labels)
    for key, real in real_cells.items():
        assert labels[key] == real, key
        assert labels[key].encode("utf-8") == real.encode("utf-8"), key


def test_b3_report_meta_fills_header_fields():
    tree = build_tree(REAL_PHENOMENON_TORN_CASE)
    md = render_report(tree, meta={
        "report_id": "KTD-2026-01-0001",
        "author": "Nguyen Van A",
        "date": "2026-10-01",
        "model": "Iris2024",
        "line": "C33-A1",
    })
    assert "KTD-2026-01-0001" in md
    assert "Nguyen Van A" in md
    assert "2026-10-01" in md
    assert "Iris2024" in md
    assert "C33-A1" in md


def test_b3_export_report_md_dispatch(tmp_path):
    tree = build_tree(REAL_PHENOMENON_STOPPER)
    out = export_report(tree, tmp_path / "plan.md")
    assert out.suffix == ".md" and out.is_file()
    assert out.read_text(encoding="utf-8") == render_report(tree)
    # Unknown suffix falls back to .md.
    out2 = export_report(tree, tmp_path / "plan.txt")
    assert out2.suffix == ".md" and out2.is_file()


def test_b3_export_docx_has_company_layout(tmp_path):
    pytest.importorskip("docx")
    tree = build_tree(REAL_PHENOMENON_CODED, code="C2340")
    out = export_docx(tree, tmp_path / "plan.docx",
                      meta={"report_id": "KTD-2026-01-0002"})
    assert out.suffix == ".docx" and out.is_file()

    from docx import Document
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any(REPORT_TITLE in t for t in texts)
    assert any("KTD-2026-01-0002" in t for t in texts)
    assert any("Cây điều tra 4M" in t for t in texts)
    assert any("Chuỗi Why-Why" in t for t in texts)
    # The Item/Details table mirrors the KTD sheet (12 detail rows).
    table = doc.tables[0]
    assert table.cell(0, 0).text == "Item／項目"
    assert table.cell(0, 1).text == "Details／詳細"
    assert len(table.rows) == 1 + len(REPORT_FIELDS) - 1  # minus investigation
    flat = " ".join(c.text for row in table.rows for c in row.cells)
    assert "Contents of defect" in flat
    assert REAL_PHENOMENON_CODED in flat
    # Two-line KTD labels keep a real line break inside the docx cell
    # (byte-exact with the form), not a flattened "／".
    label_cells = [table.cell(r, 0).text for r in range(1, len(table.rows))]
    assert "Status of occurrence at Line\nラインでの発生状況" in label_cells
    assert "Machine\u3000No.／仕上げ-マシンNo." in label_cells


def test_b3_export_report_docx_dispatch(tmp_path):
    pytest.importorskip("docx")
    tree = build_tree(REAL_PHENOMENON_TORN_CASE)
    out = export_report(tree, tmp_path / "plan.docx")
    assert out.suffix == ".docx" and out.is_file()


# ---------------------------------------------------------------------------
# Chat action wiring
# ---------------------------------------------------------------------------

def test_b3_chat_action_matches_and_renders():
    from aios_habit.chat_action import ChatActionRequest, match_action
    from aios_habit import chat_action_dieu_tra  # noqa: F401 - registers
    from aios_habit.chat_action import load_builtin_actions, reset_actions
    reset_actions()
    load_builtin_actions()
    req = ChatActionRequest(
        question="gợi ý hướng điều tra cho hiện tượng LCD báo C4001 khi in")
    action = match_action(req)
    assert action is not None
    assert action.name == "goi_y_huong_dieu_tra"
    outcome = action.handler(req)
    assert outcome is not None
    assert outcome.action == "goi_y_huong_dieu_tra"
    body = outcome.blocks[0].text
    assert "Cây điều tra 4M" in body
    assert "Chuỗi Why-Why" in body
    assert "LCD báo C4001" in body


def test_b3_chat_action_extracts_phenomenon():
    from aios_habit.chat_action_dieu_tra import extract_phenomenon
    assert extract_phenomenon(
        "Gợi ý hướng điều tra cho: LCD画面に2340表示") == "LCD画面に2340表示"
    assert extract_phenomenon(
        "lập cây điều tra 4M hiện tượng kẹt giấy") == "kẹt giấy"
    assert extract_phenomenon("gợi ý hướng điều tra") == ""


def test_b3_chat_action_asks_when_no_phenomenon():
    from aios_habit.chat_action_dieu_tra import _handler
    from aios_habit.chat_action import ChatActionRequest
    outcome = _handler(ChatActionRequest(question="gợi ý hướng điều tra"))
    assert outcome is not None
    assert "Mô tả hiện tượng" in outcome.blocks[0].text


def test_b3_chat_action_fail_closed_on_exception(monkeypatch):
    import aios_habit.chat_action_dieu_tra as mod
    from aios_habit.chat_action import ChatActionRequest

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(mod, "build_tree", boom)
    outcome = mod._handler(
        ChatActionRequest(question="gợi ý hướng điều tra cho kẹt giấy"))
    assert outcome is None  # chat falls back to the normal RAG flow


def test_b3_chat_action_fail_closed_on_bad_input():
    from aios_habit.chat_action_dieu_tra import _handler
    from aios_habit.chat_action import ChatActionRequest
    # None/empty question must not raise; the action asks for the phenomenon.
    outcome = _handler(ChatActionRequest(question=None))
    assert outcome is not None
    assert "Mô tả hiện tượng" in outcome.blocks[0].text
