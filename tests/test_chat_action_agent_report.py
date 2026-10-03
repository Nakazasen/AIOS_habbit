"""Tests cho chat action tao/sua bao cao bang loi (UX-AGENT-REPORT items 1-2).

Bao phu: phan tich lenh (tao/sua/dieu-tra-nhuong/clarify), chay dau-cuoi qua
ChatActionRequest voi backup bat buoc, the dinh kem file trong cau tra loi,
uu tien action bao cao dieu tra, va vong lap feedback + metric.
"""

import json

import pytest

from aios_habit import agent_report_feedback as feedback
from aios_habit import chat_action as chat_action_framework
from aios_habit import chat_action_agent_report as report_action
from aios_habit.chat_action import (
    ChatActionRequest,
    load_builtin_actions,
    match_action,
    match_all_actions,
)


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("AIOS_DOC_ROOT", str(tmp_path / "bao_cao"))
    monkeypatch.setenv("AIOS_LOCAL_CASES_DIR", str(tmp_path / "local_cases"))
    monkeypatch.chdir(tmp_path)
    yield tmp_path


def _request(question, **context):
    return ChatActionRequest(question=question, context=dict(context))


def _artifact_meta(outcome):
    text = "\n".join(block.text for block in outcome.blocks)
    marker = "aios_chat_artifact:"
    assert marker in text
    start = text.index(marker) + len(marker)
    end = text.index("-->", start)
    return json.loads(text[start:end].strip())


# --- Phan tich lenh ---


def test_parse_create_with_filename_and_content():
    cmd = report_action.parse_report_command(
        "tạo báo cáo tuan.md: tuần này line 1 chạy ổn định"
    )
    assert cmd.kind == "create"
    assert cmd.filename == "tuan.md"
    assert cmd.operations


def test_parse_create_without_filename_auto_names_md():
    cmd = report_action.parse_report_command("tạo báo cáo tuần cho line 1")
    assert cmd.kind == "create"
    assert cmd.filename.endswith(".md")
    assert cmd.filename.startswith("bao-cao-")


def test_parse_create_slide_keywords_make_pptx():
    cmd = report_action.parse_report_command("tạo slide báo cáo tuần")
    assert cmd.kind == "create"
    assert cmd.filename.endswith(".pptx")


def test_parse_edit_replace():
    cmd = report_action.parse_report_command(
        "sửa bc.md: thay số lỗi 5 thành số lỗi 3"
    )
    assert cmd.kind == "edit"
    assert cmd.filename == "bc.md"
    assert cmd.operations[0]["op"] == "replace"
    assert cmd.operations[0]["old"] == "số lỗi 5"
    assert cmd.operations[0]["new"] == "số lỗi 3"


def test_parse_edit_append_to_slide_number():
    cmd = report_action.parse_report_command(
        "sửa sl.pptx: thêm vào slide 2: bổ sung số liệu tháng 10"
    )
    assert cmd.kind == "edit"
    op = cmd.operations[0]
    assert op["op"] == "append_to_slide" and op["slide"] == 2


def test_parse_edit_insert_under_heading():
    cmd = report_action.parse_report_command(
        "sửa bc.md: thêm vào mục Kết luận: line 1 đạt 98%"
    )
    assert cmd.kind == "edit"
    op = cmd.operations[0]
    assert op["op"] == "insert_under_heading"
    assert op["heading"] == "Kết luận"


def test_parse_dieu_tra_declines():
    cmd = report_action.parse_report_command("lập báo cáo điều tra cho ca F100")
    assert cmd.kind == "decline"


def test_parse_edit_without_filename_asks_clarify():
    cmd = report_action.parse_report_command("sửa báo cáo: thay X thành Y")
    assert cmd.kind == "clarify"


# --- Chay dau-cuoi ---


def test_handler_creates_md_with_attachment_card(tmp_path):
    outcome = report_action._handler(_request("tạo báo cáo tuan.md: line 1 ổn định"))
    assert outcome is not None
    target = tmp_path / "bao_cao" / "tuan.md"
    assert target.is_file()
    assert "line 1 ổn định" in target.read_text(encoding="utf-8")
    meta = _artifact_meta(outcome)
    assert meta["work_id"].startswith("ARE-")
    assert meta["result_path"] == str(target)
    assert meta["work_type"] == "agent_report_edit"


def test_handler_edits_with_backup(tmp_path):
    root = tmp_path / "bao_cao"
    root.mkdir(parents=True, exist_ok=True)
    target = root / "bc.md"
    target.write_text("# Cũ\nNội dung gốc.\n", encoding="utf-8")
    outcome = report_action._handler(
        _request("sửa bc.md: thay Nội dung gốc. thành Nội dung mới.")
    )
    assert outcome is not None
    text = target.read_text(encoding="utf-8")
    assert "Nội dung mới." in text
    meta = _artifact_meta(outcome)
    backup = tmp_path / "bao_cao" / meta["checkpoint_path"].split("/")[-1]
    assert backup.is_file()
    assert "Nội dung gốc." in backup.read_text(encoding="utf-8")


def test_handler_edit_missing_file_explains(tmp_path):
    outcome = report_action._handler(_request("sửa khong-co.md: thêm đoạn mới"))
    assert outcome is not None
    text = "\n".join(block.text for block in outcome.blocks)
    assert "Chưa tìm thấy" in text
    assert "aios_chat_artifact" not in text


def test_handler_unsupported_suffix_refused():
    outcome = report_action._handler(_request("tạo báo cáo a.exe: nội dung"))
    assert outcome is not None
    text = "\n".join(block.text for block in outcome.blocks)
    assert ".docx / .pptx / .md" in text


def test_handler_docx_create_and_edit(tmp_path):
    pytest.importorskip("docx")
    created = report_action._handler(_request("tạo báo cáo bc.docx: dòng mở đầu"))
    assert created is not None
    target = tmp_path / "bao_cao" / "bc.docx"
    edited = report_action._handler(
        _request("sửa bc.docx: thay dòng mở đầu thành dòng đã sửa")
    )
    assert edited is not None
    from docx import Document

    texts = [p.text for p in Document(str(target)).paragraphs]
    assert any("dòng đã sửa" in t for t in texts)
    meta = _artifact_meta(edited)
    assert meta["checkpoint_path"]


def test_handler_pptx_slide_targeted_edit(tmp_path):
    pytest.importorskip("pptx")
    created = report_action._handler(
        _request("tạo báo cáo trình chiếu sl.pptx: tổng quan")
    )
    assert created is not None
    edited = report_action._handler(
        _request("sửa sl.pptx: thêm vào slide 1: bổ sung số liệu tháng 10")
    )
    assert edited is not None
    from pptx import Presentation

    prs = Presentation(str(tmp_path / "bao_cao" / "sl.pptx"))
    all_text = " ".join(
        shape.text for slide in prs.slides for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "bổ sung số liệu tháng 10" in all_text


def test_handler_dieu_tra_yields_to_specialist_action():
    outcome = report_action._handler(_request("lập báo cáo điều tra cho ca F100"))
    assert outcome is None


def test_investigation_report_keeps_priority_over_mine():
    # "tra_cuu_loi_tuong_tu" dang ky som do import vong (co san truoc ve nay);
    # dieu quan trong: trong cac action cu the (khong fallback), action bao cao
    # dieu tra chuyen biet dung truoc action tao/sua bao cao, va handler cua
    # ve nay nhuong han (tra None) de dispatch_multi bo qua.
    chat_action_framework.reset_actions()
    load_builtin_actions()
    req = _request("lập báo cáo điều tra cho ca F100")
    specific = [a for a in match_all_actions(req) if not a.fallback]
    names = [a.name for a in specific]
    assert "lap_bao_cao_dieu_tra" in names
    assert "tao_sua_bao_cao" in names
    assert names.index("lap_bao_cao_dieu_tra") < names.index("tao_sua_bao_cao")
    assert report_action._handler(req) is None


def test_action_is_registered_with_vietnamese_title():
    chat_action_framework.reset_actions()
    load_builtin_actions()
    names = {action.name for action in chat_action_framework.registered_actions()}
    assert "tao_sua_bao_cao" in names


# --- Vong lap feedback + metric ---


def test_action_log_written_on_run(tmp_path):
    report_action._handler(_request("tạo báo cáo log.md: nội dung"))
    log = tmp_path / "local_cases" / "agent_report_actions.jsonl"
    assert log.is_file()
    record = json.loads(log.read_text(encoding="utf-8").strip().splitlines()[-1])
    assert record["ok"] is True and record["filename"] == "log.md"


def test_feedback_bad_without_fields_raises():
    with pytest.raises(ValueError, match="phải ghi đủ"):
        feedback.record_report_feedback("ARE-X", "tester", "sai", ly_do="chỉ có lý do")


def test_feedback_invalid_verdict_raises():
    with pytest.raises(ValueError, match="Chấm không hợp lệ"):
        feedback.record_report_feedback("ARE-X", "tester", "tạm được")


def test_feedback_full_cycle_and_metrics(tmp_path):
    feedback.record_report_feedback("ARE-1", "tester", "dung")
    feedback.record_report_feedback(
        "ARE-2", "tester", "sai",
        ly_do="sai số", nguyen_nhan_that="nhầm file", noi_dung_nan_lai="làm lại",
    )
    metrics = feedback.feedback_metrics()
    assert metrics["total"] == 2
    assert metrics["redo_rate"] == 0.5
    actions = feedback.action_metrics()
    assert actions["total"] >= 0


def test_improvement_overview_trend_decreasing(monkeypatch):
    first = [
        {"verdict": "sai", "created_at": "2026-10-01T00:00:0%d+00:00" % i}
        for i in range(3)
    ]
    second = [
        {"verdict": "dung", "created_at": "2026-10-02T00:00:0%d+00:00" % i}
        for i in range(3)
    ]
    monkeypatch.setattr(feedback, "_read_jsonl", lambda path: first + second)
    overview = feedback.improvement_overview()
    assert overview["xu_huong"] == "giam_dan"
    assert overview["redo_rate_ky_truoc"] == 1.0
    assert overview["redo_rate_ky_nay"] == 0.0


def test_improvement_overview_not_enough_data(monkeypatch):
    monkeypatch.setattr(feedback, "_read_jsonl", lambda path: [{"verdict": "dung"}])
    assert feedback.improvement_overview()["xu_huong"] == "khong_du_du_lieu"
