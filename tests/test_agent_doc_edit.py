"""Tests cho agent_doc_edit: backup truoc ghi de, tao/sua md/docx/pptx."""

import hashlib

import pytest

from aios_habit import agent_doc_edit
from aios_habit.agent_doc_edit import backup_before_overwrite, edit_document


@pytest.fixture(autouse=True)
def _doc_root(tmp_path, monkeypatch):
    monkeypatch.setenv("AIOS_DOC_ROOT", str(tmp_path))
    monkeypatch.chdir(tmp_path)
    yield tmp_path


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_backup_matches_original(tmp_path):
    target = tmp_path / "a.md"
    target.write_text("goc", encoding="utf-8")
    result = backup_before_overwrite(target)
    assert result["ok"] is True
    assert _sha(tmp_path / result["backup"].split("/")[-1]) == _sha(target)


def test_edit_md_appends_and_replaces(tmp_path):
    target = tmp_path / "a.md"
    target.write_text("# Tieu de\nNoi dung cu.\n", encoding="utf-8")
    original_sha = _sha(target)
    result = edit_document(
        target,
        [
            {"op": "replace", "old": "Noi dung cu.", "new": "Noi dung moi."},
            {"op": "append_text", "text": "Them dong."},
        ],
    )
    assert result["ok"] is True
    assert result["backup"] is not None
    assert _sha(tmp_path / result["backup"].split("/")[-1]) == original_sha
    text = target.read_text(encoding="utf-8")
    assert "Noi dung moi." in text and "Them dong." in text


def test_create_md_when_missing(tmp_path):
    target = tmp_path / "moi.md"
    result = edit_document(target, [{"op": "append_text", "text": "Xin chao"}], create=True)
    assert result["ok"] is True
    assert result["backup"] is None
    assert "Xin chao" in target.read_text(encoding="utf-8")


def test_edit_missing_file_fails(tmp_path):
    result = edit_document(tmp_path / "khong-co.md", [{"op": "append_text", "text": "x"}])
    assert result["ok"] is False


def test_replace_missing_text_fails_and_keeps_backup(tmp_path):
    target = tmp_path / "a.md"
    target.write_text("abc", encoding="utf-8")
    result = edit_document(target, [{"op": "replace", "old": "zzz", "new": "q"}])
    assert result["ok"] is False
    assert "Không tìm thấy" in result["error_vi"]
    assert target.read_text(encoding="utf-8") == "abc"


def test_path_traversal_refused(tmp_path):
    result = edit_document("../../etc/passwd", [{"op": "append_text", "text": "x"}], create=True)
    assert result["ok"] is False
    assert "ngoài thư mục" in result["error_vi"]


def test_unsupported_suffix_refused(tmp_path):
    result = edit_document(tmp_path / "a.exe", [{"op": "append_text", "text": "x"}], create=True)
    assert result["ok"] is False


def test_docx_create_and_edit(tmp_path):
    pytest.importorskip("docx")
    target = tmp_path / "bc.docx"
    created = edit_document(
        target, [{"op": "append_paragraph", "text": "Dong 1"}], create=True
    )
    assert created["ok"] is True
    edited = edit_document(
        target,
        [
            {"op": "replace", "old": "Dong 1", "new": "Dong 1 sua"},
            {"op": "append_paragraph", "text": "Dong 2"},
        ],
    )
    assert edited["ok"] is True and edited["backup"] is not None
    from docx import Document
    texts = [p.text for p in Document(str(target)).paragraphs]
    assert any("Dong 1 sua" in t for t in texts)
    assert any("Dong 2" in t for t in texts)


def test_pptx_create_append_slide(tmp_path):
    pytest.importorskip("pptx")
    target = tmp_path / "sl.pptx"
    result = edit_document(
        target,
        [{"op": "append_slide", "title": "Tieu de", "bullets": ["y 1", "y 2"]}],
        create=True,
    )
    assert result["ok"] is True
    from pptx import Presentation
    prs = Presentation(str(target))
    assert len(prs.slides) == 1
    all_text = " ".join(
        shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Tieu de" in all_text and "y 1" in all_text


def test_pptx_backup_before_overwrite(tmp_path):
    pytest.importorskip("pptx")
    target = tmp_path / "sl.pptx"
    edit_document(target, [{"op": "append_slide", "title": "A"}], create=True)
    original_sha = _sha(target)
    result = edit_document(target, [{"op": "append_slide", "title": "B"}])
    assert result["ok"] is True
    assert _sha(tmp_path / result["backup"].split("/")[-1]) == original_sha
