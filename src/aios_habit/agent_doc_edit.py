"""Sua tao tai lieu docx/pptx/md bang lenh loi — lop engine an toan (UX-AGENT-REPORT).

Nguyen tac:
- Moi ghi de file da ton tai BAT BUOC backup truoc; backup phai khop SHA-256
  voi ban goc thi moi cho ghi tiep.
- Chi cho phep trong goc lam viec (AIOS_DOC_ROOT hoac thu muc hien tai) —
  chan path traversal.
- Hieu NL -> operations do lane LLM dam nhan luc chay; day la lop thuc thi
  xac dinh, test duoc tren VM.

Operations:
- {"op": "append_text", "text": "..."}            (md/txt: noi vao cuoi)
- {"op": "replace", "old": "...", "new": "...", "count": 1}
- {"op": "append_paragraph", "text": "..."}        (docx)
- {"op": "append_slide", "title": "...", "bullets": ["..."]}  (pptx)

Tuong thich Python 3.11.
"""

from __future__ import annotations

import hashlib
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

ALLOWED_SUFFIXES = {".md", ".txt", ".docx", ".pptx"}
_BACKUP_SUFFIX = ".bak"


def _doc_root() -> Path:
    root = os.environ.get("AIOS_DOC_ROOT", "")
    return Path(root).resolve() if root else Path.cwd().resolve()


def _safe_path(path: str | Path) -> Path:
    """Resolve va chan path traversal ra ngoai goc lam viec."""
    candidate = (Path(path) if Path(path).is_absolute() else _doc_root() / path).resolve()
    root = _doc_root()
    if candidate != root and root not in candidate.parents:
        raise ValueError("Đường dẫn nằm ngoài thư mục làm việc, từ chối.")
    return candidate


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def backup_before_overwrite(path: str | Path) -> Dict[str, Any]:
    """Backup file truoc khi ghi de. Tra ve {"ok", "backup", "sha256"}."""
    target = _safe_path(path)
    if not target.is_file():
        return {"ok": False, "error_vi": f"Không tìm thấy file để backup: {target.name}."}
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    backup = target.with_name(f"{target.name}{_BACKUP_SUFFIX}-{stamp}")
    counter = 0
    while backup.exists():
        counter += 1
        backup = target.with_name(f"{target.name}{_BACKUP_SUFFIX}-{stamp}-{counter}")
    backup.write_bytes(target.read_bytes())
    if _sha256(backup) != _sha256(target):
        backup.unlink(missing_ok=True)
        return {"ok": False, "error_vi": "Backup không khớp bản gốc, dừng để an toàn."}
    return {"ok": True, "backup": str(backup), "sha256": _sha256(target)}


def _check_suffix(target: Path, for_create: bool) -> Dict[str, Any] | None:
    if target.suffix.lower() not in ALLOWED_SUFFIXES:
        return {
            "ok": False,
            "error_vi": f"Chỉ hỗ trợ {sorted(ALLOWED_SUFFIXES)}, file này là {target.suffix or 'không rõ'}.",
        }
    if not for_create and not target.is_file():
        return {"ok": False, "error_vi": f"Không tìm thấy file để sửa: {target.name}."}
    return None


def _apply_text_ops(target: Path, operations: List[Dict[str, Any]]) -> int:
    text = target.read_text(encoding="utf-8") if target.exists() else ""
    applied = 0
    for op in operations:
        kind = str(op.get("op", ""))
        if kind == "append_text":
            addition = str(op.get("text", ""))
            if not text.endswith("\n") and text:
                text += "\n"
            text += addition if addition.endswith("\n") else addition + "\n"
            applied += 1
        elif kind == "replace":
            old = str(op.get("old", ""))
            new = str(op.get("new", ""))
            count = op.get("count", 1)
            if not old:
                raise ValueError("Thiếu 'old' cho thao tác replace.")
            occurrences = text.count(old)
            if occurrences == 0:
                raise ValueError(f"Không tìm thấy đoạn cần thay: {old[:60]!r}.")
            text = text.replace(old, new, int(count) if int(count) > 0 else occurrences)
            applied += 1
        else:
            raise ValueError(f"Thao tác không hỗ trợ cho text: {kind}.")
    target.write_text(text, encoding="utf-8")
    return applied


def _apply_docx_ops(target: Path, operations: List[Dict[str, Any]]) -> int:
    from docx import Document

    document = Document(str(target)) if target.exists() else Document()
    applied = 0
    for op in operations:
        kind = str(op.get("op", ""))
        if kind == "append_paragraph":
            document.add_paragraph(str(op.get("text", "")))
            applied += 1
        elif kind == "replace":
            old, new = str(op.get("old", "")), str(op.get("new", ""))
            count = int(op.get("count", 1))
            if not old:
                raise ValueError("Thiếu 'old' cho thao tác replace.")
            replaced = 0
            for para in document.paragraphs:
                if old in para.text and (count <= 0 or replaced < count):
                    for run in para.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            replaced += 1
                            break
                    else:
                        para.text = para.text.replace(old, new, 1)
                        replaced += 1
            if replaced == 0:
                raise ValueError(f"Không tìm thấy đoạn cần thay: {old[:60]!r}.")
            applied += 1
        else:
            raise ValueError(f"Thao tác không hỗ trợ cho docx: {kind}.")
    document.save(str(target))
    return applied


def _iter_text_shapes(slide: Any):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape


def _apply_pptx_ops(target: Path, operations: List[Dict[str, Any]]) -> int:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation(str(target)) if target.exists() else Presentation()
    applied = 0
    for op in operations:
        kind = str(op.get("op", ""))
        if kind == "append_slide":
            layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 else presentation.slide_layouts[0]
            slide = presentation.slides.add_slide(layout)
            title = str(op.get("title", ""))
            bullets = [str(b) for b in (op.get("bullets") or [])]
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
            else:
                box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
                box.text_frame.text = title
            if bullets:
                bodies = [s for s in _iter_text_shapes(slide) if s != slide.shapes.title]
                if bodies:
                    frame = bodies[0].text_frame
                    frame.clear()
                    for i, bullet in enumerate(bullets):
                        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                        para.text = bullet
                        para.level = 0
            applied += 1
        elif kind == "replace":
            old, new = str(op.get("old", "")), str(op.get("new", ""))
            if not old:
                raise ValueError("Thiếu 'old' cho thao tác replace.")
            replaced = 0
            for slide in presentation.slides:
                for shape in _iter_text_shapes(slide):
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                replaced += 1
            if replaced == 0:
                raise ValueError(f"Không tìm thấy đoạn cần thay: {old[:60]!r}.")
            applied += 1
        else:
            raise ValueError(f"Thao tác không hỗ trợ cho pptx: {kind}.")
    presentation.save(str(target))
    return applied


def edit_document(
    path: str | Path,
    operations: List[Dict[str, Any]],
    *,
    create: bool = False,
) -> Dict[str, Any]:
    """Thuc thi sua/tao tai lieu. Ghi de bat buoc backup truoc."""
    try:
        target = _safe_path(path)
    except ValueError as exc:
        return {"ok": False, "error_vi": str(exc)}
    err = _check_suffix(target, for_create=create)
    if err:
        return err
    if not operations:
        return {"ok": False, "error_vi": "Chưa có thao tác nào để thực hiện."}

    backup_info: Dict[str, Any] = {"backup": None}
    if target.exists():
        backup_info = backup_before_overwrite(target)
        if not backup_info.get("ok"):
            return backup_info
    try:
        suffix = target.suffix.lower()
        if suffix in {".md", ".txt"}:
            applied = _apply_text_ops(target, operations)
        elif suffix == ".docx":
            applied = _apply_docx_ops(target, operations)
        elif suffix == ".pptx":
            applied = _apply_pptx_ops(target, operations)
        else:
            return {"ok": False, "error_vi": "Định dạng không hỗ trợ."}
    except (ValueError, OSError) as exc:
        return {
            "ok": False,
            "error_vi": str(exc),
            "backup": backup_info.get("backup"),
        }
    return {
        "ok": True,
        "path": str(target),
        "backup": backup_info.get("backup"),
        "applied": applied,
        "sha256": _sha256(target),
    }
