"""Deterministic reader, normalizer, joiner, and data gate evaluator for LSU Iris datasets.

Adheres strictly to specs/008-evidence-case-loop/contracts/lsu-iris-input.md and
specs/008-evidence-case-loop/contracts/lsu-acceptance-rubric.md.
"""

from __future__ import annotations

import csv
import hashlib
import math
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

import openpyxl

from aios_habit.production_prediction.models import (
    ComponentLotMeasurement,
    DataGateReport,
    DataGateStatus,
    JigOutcomeResult,
    JoinedUnitTrace,
    LsuDatasetSnapshot,
    UnitLotLink,
)

VIETNAM_TZ = timezone(timedelta(hours=7))

# Giới hạn phần cứng laptop theo specs/008-evidence-case-loop/contracts/lsu-iris-input.md dòng 74
MAX_CSV_SIZE_BYTES = 100 * 1024 * 1024  # 100 MB
MAX_XLSX_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB
MAX_ROWS_PER_SHEET = 200_000            # 200.000 dòng
BATCH_SIZE_ROWS = 10_000                # Lô 10.000 dòng bảo vệ RAM theo T014

RECOGNIZED_UNITS = frozenset({
    "mm", "um", "nm", "cm", "m",
    "deg", "rad", "mrad",
    "mw", "w", "kw",
    "v", "mv", "kv",
    "a", "ma", "ua",
    "%", "celsius", "k",
    "s", "ms", "us", "ns", "min", "h",
    "count", "unit", "db", "hz", "khz", "mhz",
})


def compute_content_digest(path: Path) -> str:
    """Compute deterministic SHA-256 digest of file contents."""
    h = hashlib.sha256()
    with path.open("rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def _parse_time(val: Any) -> Optional[datetime]:
    """Parse time string or datetime object into Vietnam timezone. Returns None if invalid."""
    if val is None:
        return None
    if isinstance(val, datetime):
        if val.tzinfo is None:
            return val.replace(tzinfo=VIETNAM_TZ)
        return val.astimezone(VIETNAM_TZ)
    if not isinstance(val, str) or not val.strip():
        return None
    s = val.strip()
    try:
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=VIETNAM_TZ)
        else:
            dt = dt.astimezone(VIETNAM_TZ)
        return dt
    except (ValueError, TypeError):
        return None


def _iter_row_batches(path: Path, batch_size: int = BATCH_SIZE_ROWS) -> Iterator[List[Dict[str, Any]]]:
    """Yield row batches (up to batch_size rows) from CSV or XLSX file to protect laptop memory."""
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy tệp: {path.name}")
    size = path.stat().st_size
    if size == 0:
        raise ValueError(f"Tệp trống hoặc không có dữ liệu: {path.name}")

    ext = path.suffix.lower()
    if ext == ".csv" and size > MAX_CSV_SIZE_BYTES:
        size_mb = size / (1024 * 1024)
        raise ValueError(
            f"Tệp CSV vượt quá giới hạn an toàn phần cứng: {path.name} ({size_mb:.1f} MB, tối đa 100 MB)."
        )
    if ext == ".xlsx" and size > MAX_XLSX_SIZE_BYTES:
        size_mb = size / (1024 * 1024)
        raise ValueError(
            f"Tệp XLSX vượt quá giới hạn an toàn phần cứng: {path.name} ({size_mb:.1f} MB, tối đa 25 MB)."
        )

    total_rows = 0
    batch: List[Dict[str, Any]] = []

    if ext == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                headers_raw = next(reader, None)
                if not headers_raw or all(not h.strip() for h in headers_raw):
                    raise ValueError(f"Tệp CSV rỗng hoặc thiếu tiêu đề: {path.name}")
                headers = [h.strip().lower() for h in headers_raw]
                for r in reader:
                    if not r or all(not c.strip() for c in r):
                        continue
                    total_rows += 1
                    if total_rows > MAX_ROWS_PER_SHEET:
                        raise ValueError(
                            f"Tệp {path.name} vượt quá giới hạn an toàn 200.000 dòng (bảo vệ bộ nhớ laptop)."
                        )
                    row_dict = {headers[i]: r[i].strip() if i < len(r) else "" for i in range(len(headers))}
                    batch.append(row_dict)
                    if len(batch) >= batch_size:
                        yield batch
                        batch = []
                if batch:
                    yield batch
        except UnicodeDecodeError:
            raise ValueError(f"Tệp CSV không đúng chuẩn định dạng UTF-8: {path.name}")
    elif ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            sheet = wb.active
            iter_rows = sheet.iter_rows(values_only=True)
            headers_raw = next(iter_rows, None)
            if not headers_raw or all(h is None or str(h).strip() == "" for h in headers_raw):
                raise ValueError(f"Tệp Excel rỗng hoặc thiếu tiêu đề: {path.name}")
            headers = [str(h).strip().lower() if h is not None else "" for h in headers_raw]
            for r in iter_rows:
                if not r or all(c is None or str(c).strip() == "" for c in r):
                    continue
                total_rows += 1
                if total_rows > MAX_ROWS_PER_SHEET:
                    raise ValueError(
                        f"Tệp {path.name} vượt quá giới hạn an toàn 200.000 dòng (bảo vệ bộ nhớ laptop)."
                    )
                row_dict = {
                    headers[i]: str(r[i]).strip() if i < len(r) and r[i] is not None else ""
                    for i in range(len(headers))
                }
                batch.append(row_dict)
                if len(batch) >= batch_size:
                    yield batch
                    batch = []
            if batch:
                yield batch
        except ValueError:
            raise
        except Exception:
            raise ValueError(f"Không thể đọc tệp Excel {path.name}: định dạng tệp bị hỏng hoặc cấu trúc không hợp lệ.")
    else:
        raise ValueError(f"Định dạng tệp không được hỗ trợ (chỉ hỗ trợ CSV hoặc XLSX): {path.name}")

    if total_rows == 0:
        raise ValueError(f"Tệp không có dữ liệu bản ghi nào: {path.name}")


def _read_rows(path: Path) -> List[Dict[str, Any]]:
    """Read CSV or XLSX file into list of normalized column dictionaries."""
    rows: List[Dict[str, Any]] = []
    for batch in _iter_row_batches(path, batch_size=BATCH_SIZE_ROWS):
        rows.extend(batch)
    return rows


def read_lsu_source(comp_path: Path, unit_path: Path, jig_path: Path) -> LsuDatasetSnapshot:
    """Read component measurements, unit-lot links, and jig outcomes into a dataset snapshot."""
    comp_digest = compute_content_digest(comp_path)
    unit_digest = compute_content_digest(unit_path)
    jig_digest = compute_content_digest(jig_path)

    now = datetime.now(VIETNAM_TZ)
    parse_errors: List[str] = []

    comp_list: List[ComponentLotMeasurement] = []
    comp_idx = 0
    for batch in _iter_row_batches(comp_path, batch_size=BATCH_SIZE_ROWS):
        for r in batch:
            comp_idx += 1
            m_id = r.get("lot_measurement_id") or f"{comp_digest[:8]}_m_{comp_idx:05d}"
            val_raw = r.get("value")
            val: Optional[float] = None
            if val_raw is not None and str(val_raw).strip() != "":
                try:
                    val = float(val_raw)
                    if math.isnan(val) or math.isinf(val):
                        val = None
                        parse_errors.append(f"Tệp {comp_path.name} dòng {comp_idx}: giá trị đo là NaN hoặc Inf.")
                except (ValueError, TypeError):
                    val = None
                    parse_errors.append(f"Tệp {comp_path.name} dòng {comp_idx}: giá trị đo '{val_raw}' không phải số thực hợp lệ.")
            else:
                parse_errors.append(f"Tệp {comp_path.name} dòng {comp_idx}: thiếu giá trị đo.")

            time_raw = r.get("event_time", "")
            ev_time = _parse_time(time_raw)
            if ev_time is None:
                parse_errors.append(f"Tệp {comp_path.name} dòng {comp_idx}: thời điểm đo '{time_raw}' không đúng định dạng thời gian.")

            comp_list.append(
                ComponentLotMeasurement(
                    lot_measurement_id=m_id,
                    component_lot_id=r.get("component_lot_id", ""),
                    component_code=r.get("component_code", ""),
                    metric_name=r.get("metric_name", ""),
                    value=val,
                    unit=r.get("unit", ""),
                    event_time=ev_time,
                    ingested_at=now,
                    source_digest=comp_digest,
                )
            )

    unit_list: List[UnitLotLink] = []
    unit_idx = 0
    for batch in _iter_row_batches(unit_path, batch_size=BATCH_SIZE_ROWS):
        for r in batch:
            unit_idx += 1
            l_id = r.get("link_id") or f"{unit_digest[:8]}_l_{unit_idx:05d}"
            time_raw = r.get("assembly_time", "")
            as_time = _parse_time(time_raw)
            if as_time is None:
                parse_errors.append(f"Tệp {unit_path.name} dòng {unit_idx}: thời gian lắp ráp '{time_raw}' không đúng định dạng thời gian.")

            unit_list.append(
                UnitLotLink(
                    link_id=l_id,
                    unit_serial=r.get("unit_serial", ""),
                    component_lot_id=r.get("component_lot_id", ""),
                    component_code=r.get("component_code", ""),
                    assembly_time=as_time,
                    line_id=r.get("line_id"),
                    station_id=r.get("station_id"),
                    ingested_at=now,
                    source_digest=unit_digest,
                )
            )

    jig_list: List[JigOutcomeResult] = []
    jig_idx = 0
    for batch in _iter_row_batches(jig_path, batch_size=BATCH_SIZE_ROWS):
        for r in batch:
            jig_idx += 1
            j_id = r.get("jig_result_id") or f"{jig_digest[:8]}_j_{jig_idx:05d}"
            val_raw = r.get("value")
            val = None
            if val_raw is not None and str(val_raw).strip() != "":
                try:
                    val = float(val_raw)
                    if math.isnan(val) or math.isinf(val):
                        val = None
                        parse_errors.append(f"Tệp {jig_path.name} dòng {jig_idx}: giá trị JIG là NaN hoặc Inf.")
                except (ValueError, TypeError):
                    val = None
                    parse_errors.append(f"Tệp {jig_path.name} dòng {jig_idx}: giá trị JIG '{val_raw}' không phải số thực hợp lệ.")

            time_raw = r.get("event_time", "")
            ev_time = _parse_time(time_raw)
            if ev_time is None:
                parse_errors.append(f"Tệp {jig_path.name} dòng {jig_idx}: thời điểm JIG '{time_raw}' không đúng định dạng thời gian.")

            jig_list.append(
                JigOutcomeResult(
                    jig_result_id=j_id,
                    unit_serial=r.get("unit_serial", ""),
                    jig_id=r.get("jig_id", ""),
                    run_id=r.get("run_id", ""),
                    event_time=ev_time,
                    metric_name=r.get("metric_name", ""),
                    value=val,
                    unit=r.get("unit") or None,
                    jig_version=r.get("jig_version", ""),
                    process_version=r.get("process_version", ""),
                    target_label=r.get("target_label", "UNKNOWN").upper() or "UNKNOWN",
                    failure_code=r.get("failure_code") or None,
                    retest_outcome=r.get("retest_outcome") or None,
                    ingested_at=now,
                    source_digest=jig_digest,
                )
            )

    combined_digest = hashlib.sha256(f"{comp_digest}:{unit_digest}:{jig_digest}".encode("utf-8")).hexdigest()
    snapshot_id = f"snap_{combined_digest[:16]}"

    return LsuDatasetSnapshot(
        snapshot_id=snapshot_id,
        created_at=now,
        source_files={
            "component_lots": comp_path.name,
            "unit_lots": unit_path.name,
            "jig_outcomes": jig_path.name,
        },
        component_measurements=comp_list,
        unit_links=unit_list,
        jig_outcomes=jig_list,
        content_digest=combined_digest,
        parse_errors=parse_errors,
    )


def normalize_records(snapshot: LsuDatasetSnapshot) -> LsuDatasetSnapshot:
    """Normalize timestamps to Vietnam timezone, float numbers, and strip units."""
    norm_comp = []
    for m in snapshot.component_measurements:
        ev = None
        if m.event_time is not None:
            ev = m.event_time.astimezone(VIETNAM_TZ) if m.event_time.tzinfo else m.event_time.replace(tzinfo=VIETNAM_TZ)
        v = float(m.value) if m.value is not None else None
        norm_comp.append(
            ComponentLotMeasurement(
                lot_measurement_id=m.lot_measurement_id.strip(),
                component_lot_id=m.component_lot_id.strip(),
                component_code=m.component_code.strip(),
                metric_name=m.metric_name.strip(),
                value=v,
                unit=m.unit.strip(),
                event_time=ev,
                ingested_at=m.ingested_at,
                source_digest=m.source_digest,
            )
        )

    norm_units = []
    for u in snapshot.unit_links:
        as_t = None
        if u.assembly_time is not None:
            as_t = u.assembly_time.astimezone(VIETNAM_TZ) if u.assembly_time.tzinfo else u.assembly_time.replace(tzinfo=VIETNAM_TZ)
        norm_units.append(
            UnitLotLink(
                link_id=u.link_id.strip(),
                unit_serial=u.unit_serial.strip(),
                component_lot_id=u.component_lot_id.strip(),
                component_code=u.component_code.strip(),
                assembly_time=as_t,
                line_id=u.line_id.strip() if u.line_id else None,
                station_id=u.station_id.strip() if u.station_id else None,
                ingested_at=u.ingested_at,
                source_digest=u.source_digest,
            )
        )

    norm_jigs = []
    for j in snapshot.jig_outcomes:
        ev = None
        if j.event_time is not None:
            ev = j.event_time.astimezone(VIETNAM_TZ) if j.event_time.tzinfo else j.event_time.replace(tzinfo=VIETNAM_TZ)
        norm_jigs.append(
            JigOutcomeResult(
                jig_result_id=j.jig_result_id.strip(),
                unit_serial=j.unit_serial.strip(),
                jig_id=j.jig_id.strip(),
                run_id=j.run_id.strip(),
                event_time=ev,
                metric_name=j.metric_name.strip(),
                value=float(j.value) if j.value is not None else None,
                unit=j.unit.strip() if j.unit else None,
                jig_version=j.jig_version.strip(),
                process_version=j.process_version.strip(),
                target_label=j.target_label.strip().upper(),
                failure_code=j.failure_code.strip() if j.failure_code else None,
                retest_outcome=j.retest_outcome.strip() if j.retest_outcome else None,
                ingested_at=j.ingested_at,
                source_digest=j.source_digest,
            )
        )

    return LsuDatasetSnapshot(
        snapshot_id=snapshot.snapshot_id,
        created_at=snapshot.created_at,
        source_files=snapshot.source_files,
        component_measurements=norm_comp,
        unit_links=norm_units,
        jig_outcomes=norm_jigs,
        content_digest=snapshot.content_digest,
        parse_errors=list(snapshot.parse_errors),
    )


def validate_records(snapshot: LsuDatasetSnapshot) -> List[str]:
    """Validate presence of required key fields across records."""
    errors = []
    for m in snapshot.component_measurements:
        if not m.component_lot_id:
            errors.append(f"Bản ghi đo {m.lot_measurement_id} thiếu mã lô linh kiện (component_lot_id).")
    for u in snapshot.unit_links:
        if not u.unit_serial:
            errors.append(f"Liên kết {u.link_id} thiếu mã Unit (unit_serial).")
        if not u.component_lot_id:
            errors.append(f"Liên kết {u.link_id} thiếu mã lô (component_lot_id).")
    for j in snapshot.jig_outcomes:
        if not j.unit_serial:
            errors.append(f"Kết quả JIG {j.jig_result_id} thiếu mã Unit (unit_serial).")
    return errors


def join_lsu_trace(snapshot: LsuDatasetSnapshot) -> Dict[str, JoinedUnitTrace]:
    """Deterministically join lot measurements and JIG outcomes by Unit serial."""
    # Index measurements by component_lot_id
    lot_to_measurements: Dict[str, List[ComponentLotMeasurement]] = {}
    for m in snapshot.component_measurements:
        if m.component_lot_id:
            lot_to_measurements.setdefault(m.component_lot_id, []).append(m)

    # Index unit links by unit_serial
    unit_to_lots: Dict[str, List[UnitLotLink]] = {}
    for u in snapshot.unit_links:
        if u.unit_serial:
            unit_to_lots.setdefault(u.unit_serial, []).append(u)

    # Index jig outcomes by unit_serial
    unit_to_jigs: Dict[str, List[JigOutcomeResult]] = {}
    for j in snapshot.jig_outcomes:
        if j.unit_serial:
            unit_to_jigs.setdefault(j.unit_serial, []).append(j)

    traces: Dict[str, JoinedUnitTrace] = {}
    all_units = set(unit_to_lots.keys()) | set(unit_to_jigs.keys())

    for u_serial in sorted(all_units):
        links = unit_to_lots.get(u_serial, [])
        jigs = unit_to_jigs.get(u_serial, [])

        valid_links_times = [l.assembly_time for l in links if l.assembly_time is not None]
        assembly_time = min(valid_links_times, default=None)

        valid_jigs_times = [jg.event_time for jg in jigs if jg.event_time is not None]
        jig_event_time = min(valid_jigs_times, default=None)

        # Collect all component measurements for the unit's lots
        comp_measurements: List[ComponentLotMeasurement] = []
        for l in links:
            if l.component_lot_id in lot_to_measurements:
                comp_measurements.extend(lot_to_measurements[l.component_lot_id])

        # Determine target label and failure code
        target_label = "UNKNOWN"
        failure_code = None
        if jigs:
            # If any NG, unit is marked NG; otherwise first valid outcome
            ng_jigs = [jg for jg in jigs if jg.target_label == "NG"]
            if ng_jigs:
                target_label = "NG"
                failure_code = ng_jigs[0].failure_code
            elif any(jg.target_label == "OK" for jg in jigs):
                target_label = "OK"

        traces[u_serial] = JoinedUnitTrace(
            unit_serial=u_serial,
            assembly_time=assembly_time,
            jig_event_time=jig_event_time,
            target_label=target_label,
            failure_code=failure_code,
            lot_measurements=comp_measurements,
            jig_measurements=jigs,
        )

    return traces


def evaluate_data_gate_rubric(
    snapshot: LsuDatasetSnapshot,
    normalized: LsuDatasetSnapshot,
    traces: Dict[str, JoinedUnitTrace],
    target_jig_id: str = "BOWSKEW_4_BEAM",
) -> DataGateReport:
    """Evaluate dataset readiness according to LSU Iris rubric v1."""
    total_rows = len(snapshot.component_measurements) + len(snapshot.unit_links) + len(snapshot.jig_outcomes)

    # 1. Missing keys check
    validation_errors = validate_records(normalized)

    # 2. Conflicting primary keys check
    conflicting_pks = 0
    seen_comp_pks: Dict[str, Any] = {}
    for m in normalized.component_measurements:
        if m.lot_measurement_id in seen_comp_pks:
            if seen_comp_pks[m.lot_measurement_id] != m.value:
                conflicting_pks += 1
        else:
            seen_comp_pks[m.lot_measurement_id] = m.value

    seen_jig_pks: Dict[str, Any] = {}
    for j in normalized.jig_outcomes:
        k = (j.jig_result_id, j.unit_serial, j.run_id)
        if k in seen_jig_pks:
            if seen_jig_pks[k] != (j.value, j.target_label):
                conflicting_pks += 1
        else:
            seen_jig_pks[k] = (j.value, j.target_label)

    # 3. Time parse check (Rubric 2.1: ít nhất 99,5% dòng dùng để đánh giá đọc được thời gian)
    total_time_records = (
        len(normalized.component_measurements)
        + len(normalized.unit_links)
        + len(normalized.jig_outcomes)
    )
    valid_time_records = (
        sum(1 for m in normalized.component_measurements if m.event_time is not None)
        + sum(1 for u in normalized.unit_links if u.assembly_time is not None)
        + sum(1 for j in normalized.jig_outcomes if j.event_time is not None)
    )
    time_parse_valid = (valid_time_records / total_time_records * 100.0) if total_time_records > 0 else 0.0

    # 4. Corrupt measurement values check (Rubric 2.1: Phép đo phải có số thực hợp lệ)
    corrupt_val_count = sum(
        1 for m in normalized.component_measurements if m.value is None or math.isnan(m.value)
    )

    # 5. Join coverage check (Rubric 2.1: ít nhất 95% Unit nối được lot và JIG, dưới ngưỡng là BLOCKED_DATA)
    total_trace_units = len(traces)
    fully_joined_units = sum(
        1
        for t in traces.values()
        if t.lot_measurements
        and t.jig_measurements
        and t.assembly_time is not None
        and t.jig_event_time is not None
        and all(m.event_time is not None and m.value is not None for m in t.lot_measurements)
    )
    join_coverage_percent = (fully_joined_units / total_trace_units * 100.0) if total_trace_units > 0 else 0.0

    # 6. Future leakage check
    future_leaks = 0
    for t in traces.values():
        if t.jig_event_time is not None:
            for m in t.lot_measurements:
                if m.event_time is not None and m.event_time > t.jig_event_time:
                    future_leaks += 1
        if t.assembly_time is not None and t.jig_event_time is not None:
            if t.assembly_time > t.jig_event_time:
                future_leaks += 1

    # 7. Outcome counts and measurement unit recognition (Rubric 2.1 dòng 26)
    ok_count = sum(1 for t in traces.values() if t.target_label == "OK")
    ng_count = sum(1 for t in traces.values() if t.target_label == "NG")
    unknown_count = sum(1 for t in traces.values() if t.target_label == "UNKNOWN")

    total_meas_count = len(normalized.component_measurements)
    if total_meas_count > 0:
        recognized_meas_count = sum(
            1
            for m in normalized.component_measurements
            if m.unit and m.unit.strip().lower() in RECOGNIZED_UNITS
        )
        unit_recognition = (recognized_meas_count / total_meas_count) * 100.0
    else:
        unit_recognition = 100.0

    action_items: List[str] = []
    status = DataGateStatus.PASS

    if validation_errors:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(f"Thiếu trường khóa bắt buộc: {len(validation_errors)} lỗi phát hiện.")
        for err in validation_errors[:3]:
            action_items.append(f"- {err}")

    if conflicting_pks > 0:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(f"Phát hiện {conflicting_pks} khóa chính trùng nhưng mâu thuẫn dữ liệu.")

    if future_leaks > 0:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(f"Phát hiện {future_leaks} bản ghi đo xảy ra sau thời điểm JIG (rò rỉ tương lai).")

    if time_parse_valid < 99.5:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(
            f"Tỷ lệ thời gian đọc được đạt {time_parse_valid:.1f}% (dưới ngưỡng bắt buộc 99.5% theo rubric, BLOCKED_DATA)."
        )

    if unit_recognition < 99.5:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(
            f"Tỷ lệ nhận biết đơn vị đo đạt {unit_recognition:.1f}% (dưới ngưỡng bắt buộc 99.5% theo rubric, BLOCKED_DATA)."
        )

    if corrupt_val_count > 0:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(
            f"Phát hiện {corrupt_val_count} phép đo linh kiện bị hỏng hoặc thiếu số đo hợp lệ (BLOCKED_DATA)."
        )

    if snapshot.parse_errors:
        action_items.append(f"Phát hiện {len(snapshot.parse_errors)} lỗi định dạng bản ghi trong tệp nguồn.")
        for err in snapshot.parse_errors[:3]:
            action_items.append(f"- {err}")

    if join_coverage_percent < 95.0:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(
            f"Tỷ lệ nối chuỗi lot -> Unit -> JIG đạt {join_coverage_percent:.1f}% (dưới ngưỡng bắt buộc 95.0% theo rubric, BLOCKED_DATA)."
        )
    elif join_coverage_percent < 100.0:
        if status != DataGateStatus.BLOCKED_DATA:
            status = DataGateStatus.PASS_WITH_WARNING
        action_items.append(
            f"Tỷ lệ nối chuỗi lot -> Unit -> JIG đạt {join_coverage_percent:.1f}% (đạt trên 95.0% nhưng chưa đạt 100%)."
        )

    # 8. Target JIG configuration check (FR-023: Lát cắt đầu tiên phải cấu hình cho BOWSKEW 4 BEAM)
    def _normalize_jig_code(code: str) -> str:
        return code.strip().upper().replace("-", "_").replace(" ", "_")

    target_jig_norm = _normalize_jig_code(target_jig_id)
    target_jig_count = sum(
        1
        for j in normalized.jig_outcomes
        if j.jig_id and (_normalize_jig_code(j.jig_id) == target_jig_norm)
    )
    if len(normalized.jig_outcomes) > 0 and target_jig_count == 0:
        status = DataGateStatus.BLOCKED_DATA
        action_items.append(
            f"Không tìm thấy kết quả đo JIG phù hợp với đích cấu hình '{target_jig_id}' theo tiêu chuẩn FR-023 (BLOCKED_DATA)."
        )

    guidance = "Dữ liệu hợp lệ theo tiêu chuẩn LSU Iris."
    if status == DataGateStatus.BLOCKED_DATA:
        guidance = "Dữ liệu bị chặn đăng ký (BLOCKED_DATA). Cần xử lý các mục hành động trên trước khi chạy mô hình."
    elif status == DataGateStatus.PASS_WITH_WARNING:
        guidance = "Dữ liệu đạt điều kiện tối thiểu có cảnh báo (PASS_WITH_WARNING). Có thể chạy ở chế độ chỉ đọc."

    return DataGateReport(
        status=status,
        rubric_version="lsu_iris_v1",
        source_digest=snapshot.content_digest,
        total_rows=total_rows,
        join_coverage_percent=round(join_coverage_percent, 2),
        conflicting_primary_keys_count=conflicting_pks,
        future_leak_count=future_leaks,
        time_parse_valid_percent=time_parse_valid,
        unit_recognition_percent=unit_recognition,
        ok_count=ok_count,
        ng_count=ng_count,
        unknown_count=unknown_count,
        action_items=action_items,
        guidance=guidance,
    )