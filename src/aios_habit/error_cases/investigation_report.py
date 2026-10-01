"""Assemble + render a .docx investigation report for ONE error case.

This is the document builder behind the chat command "lap bao cao dieu tra
cho ca <ma phieu / ma loi>" (see chat_action_bao_cao_dieu_tra). Everything in
the report is lifted from the shipped Step 0-5 features — nothing is
re-implemented and nothing is invented:

- case facts            -> the 12 standard fields (error_cases.case_form)
- similar history       -> Step 1 engine (chat_action_error_lookup.search_similar)
- 4M tree + Why-Why     -> Step 3 (error_cases.investigation_tree.build_tree)
- auto classification   -> Step 5 (error_cases.auto_classifier.classify_error)
- recurrence alert      -> Step 5 (error_cases.auto_classifier.detect_recurrence)
- trend chart           -> the JIG SPC engine (production_prediction.spc_chart),
                           fed only with real same-code weekly counts

Anti-fabrication rule: every section that has no data says
"Chua co du lieu" explicitly. No simulated or guessed content ever lands in
the report.
"""

from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .auto_classifier import (
    ClassificationResult,
    detect_family,
    classify_error,
    detect_recurrence,
)
from .case_form import FIELDS as CASE_FIELDS
from .glossary import canonical_term
from .glossary import lookup as glossary_lookup
from .investigation_tree import InvestigationTree, build_tree
from .trend_analysis import _parse_dt

MISSING = "Chưa có dữ liệu"
MISSING_DETAIL = "Chưa có dữ liệu trong DB ca lỗi cho mục này — cần người điều tra bổ sung, agent không tự bịa."


# ---------------------------------------------------------------------------
# Small helpers over a case row (sqlite3.Row or dict)
# ---------------------------------------------------------------------------

def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _raw_value(case: Dict[str, Any], key: str) -> str:
    raw = case.get("_raw") or {}
    return _text(raw.get(key))


def case_error_code(case: Dict[str, Any]) -> str:
    """Real code if the row carries one (B0-DICT semantics): C-column first,
    then the real code extracted from column I, then the H-column group."""
    for key in ("error_code_c", "error_code_i", "error_code_h"):
        value = _text(case.get(key))
        if value:
            return value
    return ""


def case_phenomenon(case: Dict[str, Any]) -> str:
    """Phenomenon text: standard field, else raw columns I/O, else the
    investigation note — same fallback order the tree builder uses."""
    for key in ("phenomenon", "error_name"):
        value = _text(case.get(key))
        if value:
            return value
    for key in ("I", "O"):
        value = _raw_value(case, key)
        if value:
            return value
    return _text(case.get("investigation"))


def _has_column(conn: sqlite3.Connection, table: str, column: str) -> bool:
    cols = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
    return column in cols


# ---------------------------------------------------------------------------
# Data assembly
# ---------------------------------------------------------------------------

@dataclass
class InvestigationReportData:
    case: Dict[str, Any]
    code: str
    phenomenon: str
    glossary_entry: Optional[Dict[str, Any]]
    similar: List[Dict[str, Any]] = field(default_factory=list)
    countermeasures: List[Tuple[str, str]] = field(default_factory=list)
    tree: Optional[InvestigationTree] = None
    classification: Optional[ClassificationResult] = None
    recurrence: Optional[Dict[str, Any]] = None
    trend_weeks: List[Tuple[str, int]] = field(default_factory=list)
    total_same_code: int = 0
    total_cases: int = 0
    generated_on: str = ""


def _similar_cases(conn: sqlite3.Connection, case: Dict[str, Any], code: str,
                   phenomenon: str, *, top_n: int = 5) -> List[Dict[str, Any]]:
    """Reuse the Step 1 similar-case engine; drop the case itself."""
    from aios_habit.chat_action_error_lookup import search_similar

    query = f"{code} {phenomenon}".strip() or _text(case.get("no_dvd"))
    if not query:
        return []
    cards, _codes, _total = search_similar(conn, query, top_n=top_n + 2)
    own_id = case.get("id")
    return [card for card in cards if card.get("id") != own_id][:top_n]


def _countermeasures(similar: Sequence[Dict[str, Any]]) -> List[Tuple[str, str]]:
    """Countermeasures that were actually used on similar cases."""
    out: List[Tuple[str, str]] = []
    seen: set = set()
    for card in similar:
        remedy = _text(card.get("doi_sach"))
        if not remedy:
            continue
        key = remedy.casefold()
        if key in seen:
            continue
        seen.add(key)
        out.append((remedy, _text(card.get("no_dvd"))))
    return out


def _anchor_datetime(case: Dict[str, Any]) -> Optional[datetime]:
    for key in ("occurred_at", "created_at"):
        value = case.get(key)
        if value:
            try:
                return _parse_dt(value)
            except Exception:
                continue
    return None


def _recurrence(conn: sqlite3.Connection, case: Dict[str, Any], code: str) -> Optional[Dict[str, Any]]:
    """Step 5 recurrence check anchored at the case's own occurrence time.

    detect_recurrence is built for a NEW case (still absent from the DB) and
    counts it once itself (``hits + 1``). Our case is already stored, so the
    honest total including this case is ``alert.count - 1``."""
    if not code:
        return None
    anchor = _anchor_datetime(case)
    alert = detect_recurrence(conn, code, window_hours=168.0, now=anchor)
    if alert is None:
        return None
    count_incl_self = max(alert.count - 1, 0)
    if count_incl_self < 2:
        return None
    return {
        "code": code,
        "count": count_incl_self,
        "window_hours": alert.window_hours,
        "suggested_remedy": alert.suggested_remedy or "",
        "suggested_from": alert.suggested_from or "",
    }


def _trend_weeks(conn: sqlite3.Connection, code: str) -> List[Tuple[str, int]]:
    """Real weekly counts of cases carrying the same code (any code column)."""
    if not code:
        return []
    columns = ["error_code_c", "error_code_h"]
    if _has_column(conn, "error_cases", "error_code_i"):
        columns.append("error_code_i")
    where = " OR ".join(f"UPPER({col}) = UPPER(?)" for col in columns)
    sql = (
        "SELECT strftime('%Y-%W', COALESCE(occurred_at, created_at)) AS wk, "
        "COUNT(*) AS n FROM error_cases WHERE (" + where + ") "
        "GROUP BY wk ORDER BY wk"
    )
    rows = conn.execute(sql, [code] * len(columns)).fetchall()
    return [(_text(row[0]), int(row[1])) for row in rows if _text(row[0])]


def _count(conn: sqlite3.Connection, sql: str, params: Sequence[Any] = ()) -> int:
    row = conn.execute(sql, params).fetchone()
    return int(row[0]) if row and row[0] is not None else 0


def assemble_report_data(conn: sqlite3.Connection, case: Dict[str, Any]) -> InvestigationReportData:
    code = case_error_code(case)
    phenomenon = case_phenomenon(case)

    glossary_entry: Optional[Dict[str, Any]] = None
    if code:
        try:
            glossary_entry = glossary_lookup(conn, detect_family(code), code)
        except Exception:
            glossary_entry = None

    similar = _similar_cases(conn, case, code, phenomenon)
    tree = None
    if phenomenon or code:
        tree = build_tree(phenomenon, code=code, code_family=detect_family(code), conn=conn)

    classification = None
    if code or phenomenon:
        classification = classify_error(
            code,
            phenomenon=phenomenon,
            investigation=_text(case.get("investigation")),
            machine_type=_text(case.get("machine_type")) or None,
            line=_text(case.get("line")) or None,
            glossary_conn=conn,
        )

    total_same_code = 0
    if code:
        columns = ["error_code_c", "error_code_h"]
        if _has_column(conn, "error_cases", "error_code_i"):
            columns.append("error_code_i")
        where = " OR ".join(f"UPPER({col}) = UPPER(?)" for col in columns)
        total_same_code = _count(
            conn,
            "SELECT COUNT(*) FROM error_cases WHERE (" + where + ")",
            [code] * len(columns),
        )

    return InvestigationReportData(
        case=case,
        code=code,
        phenomenon=phenomenon,
        glossary_entry=glossary_entry,
        similar=similar,
        countermeasures=_countermeasures(similar),
        tree=tree,
        classification=classification,
        recurrence=_recurrence(conn, case, code),
        trend_weeks=_trend_weeks(conn, code),
        total_same_code=total_same_code,
        total_cases=_count(conn, "SELECT COUNT(*) FROM error_cases"),
        generated_on=date.today().isoformat(),
    )


# ---------------------------------------------------------------------------
# Trend chart (real data only, through the JIG SPC engine)
# ---------------------------------------------------------------------------

def render_trend_chart(data: InvestigationReportData, path: str | Path) -> Optional[Path]:
    """Weekly same-code case counts as a PNG; needs >= 3 distinct weeks."""
    if len(data.trend_weeks) < 3:
        return None
    from aios_habit.production_prediction.spc_chart import SpcChartInput, render_spc_png

    labels = [week for week, _n in data.trend_weeks]
    values = [float(n) for _week, n in data.trend_weeks]
    out = Path(path)
    chart = SpcChartInput(
        jig_id="Xu hướng lỗi",
        metric=f"Số ca mỗi tuần của mã {data.code or '(không mã)'}",
        values=values,
        nhan_thoi_gian=labels,
        nguong_tren=None,
        nguong_duoi=None,
    )
    render_spc_png(chart, out)
    return out


# ---------------------------------------------------------------------------
# .docx rendering
# ---------------------------------------------------------------------------

def _add_heading(doc: Any, text: str, level: int = 1) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _value_or_missing(value: str) -> str:
    return value if value else MISSING


def _section_case_info(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "1. Thông tin ca (12 trường chuẩn Bước 0)")
    case = data.case
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Trường"
    hdr[1].text = "Nội dung"
    for spec in CASE_FIELDS:
        label = spec["label"]
        if spec["id"] == "error_code":
            value = data.code
        else:
            column = spec.get("column") or ""
            value = _text(case.get(column)) if column else ""
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _value_or_missing(value)
    provenance = _text(case.get("source_file"))
    if provenance:
        doc.add_paragraph(
            f"Nguồn gốc ca: {provenance} › {_text(case.get('sheet_name'))} › dòng {_text(case.get('source_row')) or MISSING}."
        )


def _section_glossary(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "2. Từ điển mã lỗi")
    entry = data.glossary_entry
    if not data.code:
        doc.add_paragraph("Ca này không có mã lỗi thật trong DB. " + MISSING_DETAIL)
        return
    if not entry:
        doc.add_paragraph(f"Chưa có dữ liệu từ điển cho mã {data.code}.")
        return
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Mục"
    hdr[1].text = "Nội dung"
    for label, key in (
        ("Mã lỗi", ""),
        ("Tên gọi (Việt)", "name_vi"),
        ("Tên gọi (Nhật)", "name_ja"),
        ("Tên gọi (Anh)", "name_en"),
    ):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _value_or_missing(data.code if key == "" else _text(entry.get(key)))


def _section_similar(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "3. Ca tương tự trong lịch sử (Bước 1)")
    if not data.similar:
        doc.add_paragraph("Chưa có dữ liệu: chưa tìm thấy ca tương tự trong lịch sử. " + MISSING_DETAIL)
        return
    doc.add_paragraph(f"Đối chiếu trong tổng số {data.total_cases} ca của DB ca lỗi.")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Phiếu (No.ĐVD)", "Hiện tượng", "Nguyên nhân", "Đối sách đã áp", "Nguồn gốc"]
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for card in data.similar:
        row = table.add_row().cells
        row[0].text = _value_or_missing(_text(card.get("no_dvd")))
        row[1].text = _value_or_missing(_text(card.get("hien_tuong")))
        row[2].text = _value_or_missing(_text(card.get("nguyen_nhan")))
        row[3].text = _value_or_missing(_text(card.get("doi_sach")))
        row[4].text = _value_or_missing(_text(card.get("bao_cao_goc")))


def _section_countermeasures(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "4. Đối sách đã dùng ở ca tương tự")
    if not data.countermeasures:
        doc.add_paragraph("Chưa có dữ liệu: các ca tương tự chưa ghi đối sách. " + MISSING_DETAIL)
        return
    for remedy, no_dvd in data.countermeasures:
        label = f" (theo ca {no_dvd})" if no_dvd else ""
        doc.add_paragraph(f"- {remedy}{label}", style="List Bullet")


def _section_tree(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "5. Cây điều tra 4M + Why-Why (Bước 3)")
    tree = data.tree
    if tree is None:
        doc.add_paragraph("Chưa có dữ liệu: ca không có hiện tượng/mã để dựng cây. " + MISSING_DETAIL)
        return
    doc.add_paragraph(f"Hiện tượng: {_value_or_missing(tree.phenomenon)}")
    if tree.code:
        doc.add_paragraph(f"Mã lỗi: {tree.code} (họ {tree.code_family or '—'})")
        if tree.code_name:
            doc.add_paragraph(f"Mô tả mã: {tree.code_name}")
    for branch in tree.branches:
        doc.add_heading(f"Nhánh {branch.label}", level=2)
        for cause in branch.causes:
            doc.add_paragraph(f"- {cause}", style="List Bullet")
    _add_heading(doc, "Chuỗi Why-Why", level=2)
    if not tree.why_chain:
        doc.add_paragraph(MISSING_DETAIL)
        return
    for level, question in tree.why_chain:
        doc.add_paragraph(f"- Why {level}: {question}", style="List Bullet")
    doc.add_paragraph("Lưu ý: cây này là khung gợi ý để người điều tra xác nhận ngoài hiện trường, không phải kết luận nguyên nhân.")


def _section_classification(doc: Any, data: InvestigationReportData) -> None:
    _add_heading(doc, "6. Phân loại tự động (Bước 5 — AI gợi ý)")
    result = data.classification
    if result is None:
        doc.add_paragraph("Chưa có dữ liệu: thiếu mã/hiện tượng để phân loại. " + MISSING_DETAIL)
        return
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Mục"
    hdr[1].text = "Kết quả"
    rows = (
        ("Nhóm nguyên nhân (4M)", _text(result.cause_group_vi)),
        ("Công đoạn", _text(result.process_stage_vi)),
        ("Bộ phận phụ trách", _text(result.responsible_dept)),
        ("Độ tin cậy", f"{result.confidence:.0%}"),
        ("Phương pháp", _text(result.method)),
    )
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _value_or_missing(value)
    if result.reasons:
        _add_heading(doc, "Căn cứ phân loại", level=2)
        for reason in result.reasons:
            doc.add_paragraph(f"- {reason}", style="List Bullet")
    doc.add_paragraph("Kết quả phân loại do máy gợi ý — người điều tra xác nhận trước khi dùng chính thức.")


def _section_recurrence(doc: Any, data: InvestigationReportData, chart_path: Optional[Path]) -> None:
    _add_heading(doc, "7. Tái phát & xu hướng")
    if data.code:
        doc.add_paragraph(
            f"Tổng số ca mang mã {data.code} trong lịch sử: {data.total_same_code} ca (trong {data.total_cases} ca của DB)."
        )
    if data.recurrence:
        alert = data.recurrence
        doc.add_paragraph(
            f"CẢNH BÁO TÁI PHÁT: mã {alert['code']} đã phát sinh {alert['count']} lần "
            f"trong {alert['window_hours']:.0f} giờ quanh thời điểm ca này (tính cả ca này)."
        )
        if alert["suggested_remedy"]:
            origin = f" (từ ca {alert['suggested_from']})" if alert["suggested_from"] else ""
            doc.add_paragraph(f"Đối sách tái phát nên áp dụng ngay: {alert['suggested_remedy']}{origin}")
    elif data.code:
        doc.add_paragraph("Không phát hiện tái phát của mã này quanh thời điểm ca (theo ngưỡng Bước 5).")
    else:
        doc.add_paragraph("Chưa có dữ liệu: ca không có mã lỗi thật nên không kiểm tra được tái phát. " + MISSING_DETAIL)
    if chart_path is not None:
        doc.add_paragraph("Biểu đồ số ca cùng mã theo tuần (dữ liệu thật trong DB):")
        doc.add_picture(str(chart_path), width=_chart_width())
    else:
        doc.add_paragraph("Chưa đủ dữ liệu để vẽ biểu đồ xu hướng (cần ca cùng mã ở ít nhất 3 tuần khác nhau).")


def _chart_width() -> Any:
    from docx.shared import Cm

    return Cm(15)


def render_report_docx(data: InvestigationReportData, out_path: str | Path,
                       chart_path: Optional[Path] = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    case = data.case
    no_dvd = _text(case.get("no_dvd")) or "(chưa có số phiếu)"

    title = doc.add_heading(f"Báo cáo điều tra lỗi/調査報告書 — {no_dvd}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(f"Người lập báo cáo (tác giả): ………………………   Ngày lập: {data.generated_on}")
    phenomenon_line = _value_or_missing(data.phenomenon)
    code_line = data.code or MISSING
    doc.add_paragraph(f"Hiện tượng: {phenomenon_line}   |   Mã lỗi: {code_line}")

    _section_case_info(doc, data)
    _section_glossary(doc, data)
    _section_similar(doc, data)
    _section_countermeasures(doc, data)
    _section_tree(doc, data)
    _section_classification(doc, data)
    _section_recurrence(doc, data, chart_path)

    doc.add_paragraph(
        "— Báo cáo do agent tự lắp từ DB ca lỗi Bước 0–5 (thông tin ca, lịch sử ca tương tự, "
        "cây 4M + Why-Why, phân loại và cảnh báo tái phát tự động). Mục nào không có dữ liệu đều "
        "ghi rõ 'Chưa có dữ liệu'; agent không tự bịa số liệu. Cần người điều tra xác nhận trước khi dùng chính thức."
    )
    doc.save(str(out))
    return out


# ---------------------------------------------------------------------------
# One-call builder
# ---------------------------------------------------------------------------

_SLUG_BAD = re.compile(r"[^A-Za-z0-9._-]+")


def case_file_slug(case: Dict[str, Any]) -> str:
    no_dvd = _text(case.get("no_dvd")) or "khong-so-phieu"
    return _SLUG_BAD.sub("_", no_dvd).strip("_") or "khong-so-phieu"


def build_investigation_report(
    conn: sqlite3.Connection,
    case: Dict[str, Any],
    out_dir: str | Path,
    *,
    stamp: Optional[str] = None,
) -> Path:
    """Assemble -> chart -> .docx; returns the .docx path."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    data = assemble_report_data(conn, case)
    stamp = stamp or datetime.now().strftime("%Y%m%d-%H%M%S")
    slug = case_file_slug(case)
    chart_path = render_trend_chart(data, out / f"xu-huong-{slug}-{stamp}.png")
    docx_path = out / f"bao-cao-dieu-tra-{slug}-{stamp}.docx"
    render_report_docx(data, docx_path, chart_path=chart_path)
    return docx_path
