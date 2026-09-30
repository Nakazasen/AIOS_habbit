"""F2: importer for the 29-column "History KDTPS" sheet -> error_cases.

Source: Loi KDTPS.xlsx, sheet "History KDTPS" (header row 4, data from
row 5, 29 columns A-AC). The legacy A-Y mapping does NOT apply here;
see column_map.normalize_history_row and recon 2026-09-27.

Rules:
- Rows missing year (A) or NO. (B) are skipped (counted).
- Rows missing machine type (D) or line (E) are skipped (the history
  dedup key needs them).
- Re-importing an unchanged file (same name + sha256 + sheet) is skipped
  whole ("vat lai file cu thi bo qua"), unless force=True.
- Green-fill detection is kept as a code path (the real file has 0
  green cells); green-skipped columns never overwrite stored values.
- ``occurred_at`` is the real occurrence date parsed from column C
  (ISO, None when the source cell is missing/unparseable); ``created_at``
  (import time) is untouched.
- Read-only: the source workbook is never modified.

``fill_occurred_at`` back-fills the column for DBs imported before it
existed (dry-run by default) without touching raw_json -- see its
docstring for why a force re-import is not used.

This module never touches the RAG index.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from . import column_map
from . import store

SHEET_NAME = "History KDTPS"
HEADER_ROW_1BASED = 4
N_COLS = 29


def _cell_rgb(cell: Any) -> Optional[Tuple[int, int, int]]:
    """Fill color of a cell as (R, G, B), or None when unfilled."""
    fill = getattr(cell, "fill", None)
    if fill is None or not getattr(fill, "patternType", None):
        return None
    color = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
    rgb = getattr(color, "rgb", None)
    if not rgb or rgb == "00000000":
        return None
    hex_rgb = str(rgb)
    if len(hex_rgb) == 8:  # ARGB -> drop alpha
        hex_rgb = hex_rgb[2:]
    if len(hex_rgb) != 6:
        return None
    return (int(hex_rgb[0:2], 16), int(hex_rgb[2:4], 16), int(hex_rgb[4:6], 16))


def green_skip_letters(cells: List[Any]) -> List[str]:
    """Excel column letters whose fill is the legacy 'do not overwrite' green."""
    return [
        column_map.col_letter(i)
        for i, cell in enumerate(cells)
        if column_map.is_green_skip(_cell_rgb(cell))
    ]


def find_existing_batch(
    conn: sqlite3.Connection, source_file: str, file_sha256: str, sheet_name: str
) -> Optional[int]:
    row = conn.execute(
        """SELECT id FROM import_batches
           WHERE source_file = ? AND file_sha256 = ? AND sheet_name = ?""",
        (source_file, file_sha256, sheet_name),
    ).fetchone()
    return int(row["id"]) if row else None


def _is_empty_row(values: List[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def _valid_no_dvd(no_dvd: str) -> bool:
    """False when the year (A) or NO. (B) side is missing: '', '/', '/5', '2024/'."""
    return bool(no_dvd) and no_dvd != "/" and not no_dvd.startswith("/") \
        and not no_dvd.endswith("/")


def import_history(
    conn: sqlite3.Connection,
    xlsx_path: str | Path,
    *,
    sheet_name: str = SHEET_NAME,
    force: bool = False,
) -> Dict[str, Any]:
    """Import the History KDTPS sheet. Returns a result dict.

    status is "imported" or "skipped" (unchanged file re-import).
    """
    path = Path(xlsx_path)
    sha = store.sha256_file(path)
    source_file = path.name

    if not force:
        existing = find_existing_batch(conn, source_file, sha, sheet_name)
        if existing is not None:
            return {
                "status": "skipped",
                "batch_id": existing,
                "rows_read": 0,
                "rows_imported": 0,
                "rows_skipped": 0,
            }
        batch_id = store.start_batch(
            conn,
            source_file=source_file,
            file_sha256=sha,
            sheet_name=sheet_name,
            sheet_type=None,  # history sheet has no Máy in/KIT distinction
            header_row=HEADER_ROW_1BASED,
            notes="F2: 29-column History KDTPS format (format=history_29)",
        )
    else:
        # Refresh: re-run into the existing batch so provenance stays one row
        # per (file, sha, sheet); upserts update the cases in place.
        batch_id = find_existing_batch(conn, source_file, sha, sheet_name)
        if batch_id is None:
            batch_id = store.start_batch(
                conn,
                source_file=source_file,
                file_sha256=sha,
                sheet_name=sheet_name,
                sheet_type=None,
                header_row=HEADER_ROW_1BASED,
                notes="F2: 29-column History KDTPS format (format=history_29)",
            )

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        rows_read = rows_imported = rows_skipped = rows_with_occurred_at = 0
        for excel_row, row_cells in enumerate(
            ws.iter_rows(min_row=HEADER_ROW_1BASED + 1, max_col=N_COLS),
            start=HEADER_ROW_1BASED + 1,
        ):
            values = [c.value for c in row_cells]
            if _is_empty_row(values):
                continue
            rows_read += 1
            fields = column_map.normalize_history_row(values)
            no_dvd = (fields.get("no_dvd") or "").strip()
            if not _valid_no_dvd(no_dvd):
                rows_skipped += 1  # missing year (A) or NO. (B)
                continue
            try:
                store.upsert_case(
                    conn,
                    batch_id=batch_id,
                    source_row=excel_row,
                    fields=fields,
                    skip_cells=green_skip_letters(row_cells),
                    format="history_29",
                )
            except ValueError:
                rows_skipped += 1  # missing machine (D) / line (E)
                continue
            if fields.get("occurred_at"):
                rows_with_occurred_at += 1
            rows_imported += 1
    finally:
        wb.close()

    store.finish_batch(
        conn,
        batch_id,
        rows_read=rows_read,
        rows_imported=rows_imported,
        rows_skipped=rows_skipped,
    )
    return {
        "status": "imported",
        "batch_id": batch_id,
        "rows_read": rows_read,
        "rows_imported": rows_imported,
        "rows_skipped": rows_skipped,
        "rows_with_occurred_at": rows_with_occurred_at,
    }


def fill_occurred_at(
    conn: sqlite3.Connection,
    xlsx_path: str | Path,
    *,
    sheet_name: str = SHEET_NAME,
    batch_id: Optional[int] = None,
    apply: bool = False,
) -> Dict[str, Any]:
    """Fill ``error_cases.occurred_at`` for rows imported before the column existed.

    Re-reads the source sheet (read-only) and updates ONLY ``occurred_at``
    (plus ``updated_at``) for the rows of one history batch, matched by
    (batch_id, source_row) and re-checked against the stored dedup key
    (no_dvd, machine_type, line); rows whose stored key no longer matches
    the source row are never touched (``mismatch``). ``raw_json`` is left
    intact on purpose: the f3b backfill keeps its provenance inside
    ``raw_json._backfill`` and a force re-import would wipe it.

    Dry-run by default (nothing written); ``apply=True`` writes and
    commits. Returns counters for the report.
    """
    path = Path(xlsx_path)
    sha = store.sha256_file(path)
    source_file = path.name
    if batch_id is None:
        batch_id = find_existing_batch(conn, source_file, sha, sheet_name)
    if batch_id is None:
        return {
            "status": "no_batch",
            "batch_id": None,
            "source_file": source_file,
            "file_sha256": sha,
            "sheet_name": sheet_name,
        }

    counters: Dict[str, Any] = {
        "rows_read": 0,
        "cases_found": 0,
        "would_update": 0,
        "updated": 0,
        "no_change": 0,
        "invalid_date": 0,
        "not_in_db": 0,
        "mismatch": 0,
    }
    samples: List[Dict[str, Any]] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for excel_row, row_cells in enumerate(
            ws.iter_rows(min_row=HEADER_ROW_1BASED + 1, max_col=N_COLS),
            start=HEADER_ROW_1BASED + 1,
        ):
            values = [c.value for c in row_cells]
            if _is_empty_row(values):
                continue
            counters["rows_read"] += 1
            fields = column_map.normalize_history_row(values)
            no_dvd = (fields.get("no_dvd") or "").strip()
            if not _valid_no_dvd(no_dvd):
                continue  # row the importer skipped -> no case
            if not (fields.get("machine_type") and fields.get("line")):
                continue  # row the importer skipped -> no case
            row = conn.execute(
                """SELECT id, occurred_at, no_dvd, machine_type, line
                   FROM error_cases WHERE batch_id = ? AND source_row = ?""",
                (batch_id, excel_row),
            ).fetchone()
            if row is None:
                counters["not_in_db"] += 1
                continue
            counters["cases_found"] += 1
            stored_key = (
                str(row["no_dvd"] or "").strip(),
                str(row["machine_type"] or "").strip(),
                str(row["line"] or "").strip(),
            )
            source_key = (
                no_dvd,
                str(fields.get("machine_type") or "").strip(),
                str(fields.get("line") or "").strip(),
            )
            if stored_key != source_key:
                counters["mismatch"] += 1
                continue
            new_value = fields.get("occurred_at")
            if not new_value:
                counters["invalid_date"] += 1
                continue
            if row["occurred_at"] == new_value:
                counters["no_change"] += 1
                continue
            counters["would_update"] += 1
            if len(samples) < 10:
                samples.append(
                    {"source_row": excel_row, "no_dvd": no_dvd, "occurred_at": new_value}
                )
            if apply:
                conn.execute(
                    """UPDATE error_cases
                       SET occurred_at = ?, updated_at = datetime('now')
                       WHERE id = ?""",
                    (new_value, row["id"]),
                )
                counters["updated"] += 1
    finally:
        wb.close()
    if apply:
        conn.commit()
    return {
        "status": "applied" if apply else "planned",
        "batch_id": batch_id,
        "source_file": source_file,
        "file_sha256": sha,
        "sheet_name": sheet_name,
        **counters,
        "samples": samples,
    }


def _main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Fill error_cases.occurred_at from the history sheet (dry-run by default)."
    )
    parser.add_argument("--db", required=True, help="SQLite path (use a copy on drive C)")
    parser.add_argument("--xlsx", required=True, help="Path to Loi KDTPS.xlsx")
    parser.add_argument("--batch-id", type=int, default=None)
    parser.add_argument("--apply", action="store_true", help="write the values")
    args = parser.parse_args(argv)
    conn = store.connect(args.db)
    try:
        result = fill_occurred_at(conn, args.xlsx, batch_id=args.batch_id, apply=args.apply)
    finally:
        conn.close()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
