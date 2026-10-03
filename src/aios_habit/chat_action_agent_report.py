"""Chat action: agent tao/sua bao cao bang loi trong khung chat (UX-AGENT-REPORT).

Backend cho item 1-2 cua ve (item 3 = phuong an UI, cho user duyet, chua code):

- `parse_report_command(question)` — hieu lenh loi dang xac dinh thanh
  ReportCommand (tao moi / sua file co san / can hoi lai / nhuong cho action
  chuyen biet). Dang tu do mo ("them bieu do X vao slide 3") do lane LLM dam
  nhan luc chay; tang nay xu ly dang dien dat ro rang va phuc vu test tren VM.
- `run_report_edit(...)` — thuc thi tao/sua qua `agent_doc_edit` (backup truoc
  moi ghi de, chan path traversal), ghi log van hanh + tra the dinh kem file
  trong cau tra loi (metadata `aios_chat_artifact` de UI chat hien nut mo/tai).
- Vong lap cai thien: moi hanh dong ghi `local_cases/agent_report_actions.jsonl`;
  user cham dung/mot_phan/sai qua `agent_report_feedback.record_report_feedback`
  (thieu 3 truong khi che thi raise ValueError); metric + xu huong o
  `agent_report_feedback.improvement_overview`.

Fail-closed sau co `AIOS_FEATURE_CHAT_ACTION` (mac dinh tat) nhu cac action khac.
Tuong thich Python 3.11: khong f-string nhieu dong (PEP 701), khong `type` stmt.
"""

from __future__ import annotations

import json
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from aios_habit.agent_doc_edit import ALLOWED_SUFFIXES, default_doc_root, edit_document
from aios_habit.agent_report_feedback import log_report_action
from aios_habit.chat_action import (
    BLOCK_MARKDOWN,
    ChatAction,
    ChatActionBlock,
    ChatActionOutcome,
    ChatActionRequest,
    normalize_text,
    register_action,
)

ACTION_NAME = "tao_sua_bao_cao"
ACTION_TITLE = "Tạo / sửa báo cáo"

_HINTS = (
    "tao bao cao",
    "sua bao cao",
    "chinh bao cao",
    "cap nhat bao cao",
    "them vao bao cao",
    "bo sung vao bao cao",
    "lap bao cao",
    "viet bao cao",
    "tao slide",
)

# Nhuong cho action bao cao dieu tra chuyen biet (uu tien dang ky truoc).
_DECLINE_MARKERS = ("dieu tra",)

_CREATE_MARKS = ("tao", "lap", "viet")
# Sua: dong tu tran + co "bao cao" hoac ten file cu the trong cau.
_EDIT_VERB_MARKS = ("sua", "chinh", "cap nhat", "them vao", "bo sung", "thay")

_QUOTED_FILENAME_RE = re.compile(
    r"[“”\"']([\w\-\.À-ỹ\(\) ]+\.(?:docx|pptx|md|txt))[“”\"']", re.IGNORECASE
)
_FILENAME_RE = re.compile(
    r"([\w\-\.À-ỹ\(\)]+\.(?:docx|pptx|md|txt))", re.IGNORECASE
)
# Token co ve la ten file nhung duoi khong duoc ho tro (vd a.exe, b.xlsx).
_ANY_FILENAME_RE = re.compile(
    r"([\w\-\.À-ỹ\(\)]+\.[A-Za-z0-9]{1,6})\b", re.IGNORECASE
)
_REPLACE_RE = re.compile(r"thay\s+(.+?)\s+(thanh|bang)\s+(.+)", re.DOTALL)
_SLIDE_NUM_RE = re.compile(r"them vao slide\s+(\d+)\s*[:\-–]?\s*(.+)", re.DOTALL)
_SLIDE_TITLE_RE = re.compile(
    r"them vao slide\s+[“”\"'](.+?)[“”\"']\s*[:\-–]?\s*(.+)", re.DOTALL
)
_HEADING_RE = re.compile(r"them (?:vao|duoi) muc\s+(.+?)\s*[:]\s*(.+)", re.DOTALL)
_APPEND_SLIDE_RE = re.compile(r"them slide\s*[:\-–]?\s*(.+)", re.DOTALL)

_PUNCT_TAIL = " \t\r\n.,:;!?…\"'“”‘’()[]"


@dataclass
class ReportCommand:
    """Ket qua phan tich lenh loi (xac dinh, test duoc)."""

    kind: str  # "create" | "edit" | "clarify" | "decline"
    filename: str = ""
    operations: List[Dict[str, Any]] = field(default_factory=list)
    summary_vi: str = ""
    decline_reason_vi: str = ""


def _strip_quotes(text: str) -> str:
    return str(text or "").strip().strip("\"'“”‘’").strip(_PUNCT_TAIL).strip()


def _norm_flat(text: str) -> str:
    """Chuan hoa khong dau, giu nguyen do dai de map span -> raw."""
    import unicodedata

    decomposed = unicodedata.normalize("NFD", str(text or "").casefold())
    return "".join(
        ch for ch in decomposed if not unicodedata.combining(ch)
    ).replace("đ", "d")


def _raw_span(raw: str, match: "re.Match[str]", group: int) -> str:
    return _strip_quotes(raw[match.start(group):match.end(group)])


def _slugify(text: str, limit: int = 40) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", normalize_text(text)).strip("-")
    return (slug[:limit].rstrip("-")) or "bao-cao"


def _default_doc_root() -> Path:
    # Dung chung goc voi agent_doc_edit de hai lop khong bao gio tinh goc
    # khac nhau (hoi quy 2026-10-03: edit tung tu choi ~/AIOS_bao_cao).
    return default_doc_root()


def _doc_root_for(request: Optional[ChatActionRequest] = None) -> Path:
    if request is not None:
        raw = str((request.context or {}).get("doc_root") or "").strip()
        if raw:
            root = Path(raw).expanduser()
            root.mkdir(parents=True, exist_ok=True)
            return root
    return _default_doc_root()


def _pick_filename(raw: str) -> str:
    quoted = _QUOTED_FILENAME_RE.search(raw or "")
    if quoted:
        return _strip_quotes(quoted.group(1))
    match = _FILENAME_RE.search(raw or "")
    return _strip_quotes(match.group(1)) if match else ""


def _pick_unsupported_filename(raw: str) -> str:
    """Ten file dien dat ro nhung duoi khong duoc ho tro (de tu choi ro rang)."""
    for match in _ANY_FILENAME_RE.finditer(raw or ""):
        token = _strip_quotes(match.group(1))
        if Path(token).suffix.lower() not in ALLOWED_SUFFIXES:
            return token
    return ""


def _tail_after_filename(raw: str, filename: str) -> str:
    if not filename:
        return ""
    index = raw.lower().find(filename.lower())
    if index < 0:
        return ""
    return raw[index + len(filename):].strip(_PUNCT_TAIL + "-– ")


def _create_ops_for(suffix: str, title: str, body: str) -> List[Dict[str, Any]]:
    lines = [line.strip() for line in str(body or "").splitlines() if line.strip()]
    if suffix == ".pptx":
        bullets: List[str] = []
        for line in lines:
            bullets.extend(
                part.strip() for part in re.split(r"[;\n]+", line) if part.strip()
            )
        return [{"op": "append_slide", "title": title, "bullets": bullets}]
    if suffix == ".docx":
        ops: List[Dict[str, Any]] = []
        if title:
            ops.append({"op": "append_paragraph", "text": title})
        for line in lines:
            ops.append({"op": "append_paragraph", "text": line})
        return ops
    ops = []
    if title:
        ops.append({"op": "append_text", "text": "# " + title})
    for line in lines:
        ops.append({"op": "append_text", "text": line})
    return ops


def _append_ops_for(suffix: str, text: str) -> List[Dict[str, Any]]:
    cleaned = _strip_quotes(text)
    if not cleaned:
        return []
    if suffix == ".pptx":
        parts = [p.strip() for p in re.split(r"[;\n]+", cleaned) if p.strip()]
        title = parts[0][:80] if parts else cleaned[:80]
        return [{"op": "append_slide", "title": title, "bullets": parts[1:]}]
    if suffix == ".docx":
        return [{"op": "append_paragraph", "text": cleaned}]
    return [{"op": "append_text", "text": cleaned}]


def parse_report_command(question: str) -> ReportCommand:
    """Phan tich lenh loi thanh ReportCommand (khong cham file he thong)."""
    raw = str(question or "").strip()
    norm = normalize_text(raw)
    if any(marker in norm for marker in _DECLINE_MARKERS):
        return ReportCommand(
            kind="decline",
            decline_reason_vi="Lệnh về báo cáo điều tra — nhường cho action chuyên biệt.",
        )
    filename = _pick_filename(raw)
    unsupported = "" if filename else _pick_unsupported_filename(raw)
    is_create = any(mark in norm for mark in _CREATE_MARKS) and (
        "bao cao" in norm or "slide" in norm or bool(filename)
    )
    is_edit = ("bao cao" in norm or bool(filename)) and any(
        mark in norm for mark in _EDIT_VERB_MARKS
    )

    def _refuse_unsupported() -> ReportCommand:
        return ReportCommand(
            kind="clarify",
            filename=unsupported,
            summary_vi=(
                "Chỉ hỗ trợ .docx / .pptx / .md — file “%s” chưa làm được."
                % unsupported
            ),
        )

    if is_create:
        if unsupported:
            return _refuse_unsupported()
        suffix = Path(filename).suffix.lower() if filename else ""
        if not filename:
            topic = norm
            for mark in _CREATE_MARKS + ("bao cao", "slide"):
                topic = topic.replace(mark, " ")
            topic = " ".join(topic.split())
            if "slide" in norm or "trinh chieu" in norm or "powerpoint" in norm:
                suffix = ".pptx"
            elif "word" in norm:
                suffix = ".docx"
            else:
                suffix = ".md"
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
            filename = "bao-cao-%s-%s%s" % (_slugify(topic), stamp, suffix)
        title = Path(filename).stem.replace("-", " ").replace("_", " ").strip()
        body = ""
        if ":" in raw:
            body = raw.split(":", 1)[1]
        operations = _create_ops_for(suffix, title, body)
        summary = "Tạo mới %s (%s)." % (filename, suffix)
        return ReportCommand(
            kind="create", filename=filename, operations=operations, summary_vi=summary
        )

    if is_edit:
        if unsupported:
            return _refuse_unsupported()
        if not filename:
            return ReportCommand(
                kind="clarify",
                summary_vi=(
                    "Chưa rõ bạn muốn sửa file nào. Cho mình tên file "
                    "(.docx / .pptx / .md), ví dụ: “sửa báo cáo tuan.docx: "
                    "thay X thành Y”."
                ),
            )
        tail = _tail_after_filename(raw, filename)
        # Match tren ban khong dau (giu do dai) roi cat lai doan raw co dau.
        flat = _norm_flat(tail)
        operations: List[Dict[str, Any]] = []
        replace_match = _REPLACE_RE.search(flat)
        slide_num_match = _SLIDE_NUM_RE.search(flat)
        slide_title_match = _SLIDE_TITLE_RE.search(flat)
        heading_match = _HEADING_RE.search(flat)
        append_slide_match = _APPEND_SLIDE_RE.search(flat)
        suffix = Path(filename).suffix.lower()
        if replace_match:
            operations.append(
                {
                    "op": "replace",
                    "old": _raw_span(tail, replace_match, 1),
                    "new": _raw_span(tail, replace_match, 3),
                }
            )
        elif slide_num_match:
            operations.append(
                {
                    "op": "append_to_slide",
                    "slide": int(slide_num_match.group(1)),
                    "text": _raw_span(tail, slide_num_match, 2),
                }
            )
        elif slide_title_match:
            operations.append(
                {
                    "op": "append_to_slide",
                    "slide_title": _raw_span(tail, slide_title_match, 1),
                    "text": _raw_span(tail, slide_title_match, 2),
                }
            )
        elif heading_match:
            operations.append(
                {
                    "op": "insert_under_heading",
                    "heading": _raw_span(tail, heading_match, 1),
                    "text": _raw_span(tail, heading_match, 2),
                }
            )
        elif append_slide_match and suffix == ".pptx":
            operations.extend(
                _append_ops_for(suffix, _raw_span(tail, append_slide_match, 1))
            )
        elif tail:
            operations.extend(_append_ops_for(suffix, tail))
        if not operations:
            return ReportCommand(
                kind="clarify",
                filename=filename,
                summary_vi=(
                    "Chưa hiểu cần sửa gì trong %s. Bạn thử: “thay X thành Y”, "
                    "“thêm vào slide 3: …”, “thêm vào mục …: …” hoặc “thêm: …”."
                    % filename
                ),
            )
        return ReportCommand(
            kind="edit",
            filename=filename,
            operations=operations,
            summary_vi="Sửa %s (%d thao tác)." % (filename, len(operations)),
        )

    return ReportCommand(
        kind="clarify",
        summary_vi=(
            "Bạn muốn tạo mới hay sửa báo cáo nào? Ví dụ: “tạo báo cáo tuan.md: "
            "nội dung…” hoặc “sửa bao-cao.pptx: thêm slide Kết luận”."
        ),
    )


def _verify_output(path: Path) -> str:
    """Kiem tra nhe file vua ghi co mo duoc khong (khong raise ra ngoai)."""
    try:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            from docx import Document

            count = len(Document(str(path)).paragraphs)
            return "mở được, %d đoạn" % count
        if suffix == ".pptx":
            from pptx import Presentation

            count = len(Presentation(str(path)).slides)
            return "mở được, %d slide" % count
        size = path.stat().st_size
        return "mở được, %d byte" % size
    except Exception as exc:
        return "chưa xác minh được (%s)" % exc


def run_report_edit(
    filename: str,
    operations: List[Dict[str, Any]],
    *,
    create: bool = False,
    command: str = "",
    doc_root: Optional[Path] = None,
) -> Dict[str, Any]:
    """Thuc thi tao/sua bao cao; tra dict cho UI chat (khong raise)."""
    work_id = "ARE-" + uuid.uuid4().hex[:8].upper()
    root = Path(doc_root) if doc_root else _default_doc_root()
    target = root / str(filename or "").strip()
    result = edit_document(target, operations or [], create=create)
    logged = log_report_action(
        work_id=work_id,
        command=command,
        filename=str(filename),
        kind="create" if create else "edit",
        ok=bool(result.get("ok")),
        applied=int(result.get("applied") or 0),
        backup=str(result.get("backup") or ""),
        error=str(result.get("error_vi") or ""),
    )
    outcome: Dict[str, Any] = {
        "ok": bool(result.get("ok")),
        "work_id": work_id,
        "path": str(result.get("path") or target),
        "backup": str(result.get("backup") or ""),
        "applied": int(result.get("applied") or 0),
        "sha256": str(result.get("sha256") or ""),
        "error_vi": str(result.get("error_vi") or ""),
        "event_id": logged.get("event_id", ""),
    }
    if outcome["ok"]:
        outcome["verify_vi"] = _verify_output(Path(outcome["path"]))
    return outcome


def _artifact_card(result: Dict[str, Any], summary_vi: str) -> str:
    """The dinh kem file trong cau tra loi (UI chat hien nut mo/tai)."""
    path = Path(str(result.get("path", "")))
    lines = [
        "### 📄 " + (summary_vi or path.name),
        "- **File:** `%s`" % path.name,
        "- **Kiểm tra:** %s." % result.get("verify_vi", "—"),
    ]
    backup = str(result.get("backup") or "")
    if backup:
        lines.append("- **Bản gốc đã sao lưu:** `%s`" % Path(backup).name)
        lines.append("  (Muốn hoàn tác thì mở file sao lưu này để lấy lại nội dung cũ.)")
    lines.append("")
    lines.append("Mở file đính kèm bên dưới để xem ngay.")
    meta = {
        "work_id": result.get("work_id", ""),
        "work_type": "agent_report_edit",
        "result_path": str(result.get("path", "")),
        "checkpoint_path": backup,
    }
    lines.append("<!-- aios_chat_artifact: %s -->" % json.dumps(meta, ensure_ascii=False))
    return "\n".join(lines)


def _message(text: str) -> ChatActionOutcome:
    return ChatActionOutcome(
        action=ACTION_NAME,
        title=ACTION_TITLE,
        blocks=(ChatActionBlock(BLOCK_MARKDOWN, text=text),),
    )


def _handler(request: ChatActionRequest) -> Optional[ChatActionOutcome]:
    command = parse_report_command(request.question)
    if command.kind == "decline":
        return None
    if command.kind == "clarify":
        return _message(command.summary_vi)
    if command.filename and Path(command.filename).suffix.lower() not in ALLOWED_SUFFIXES:
        return _message(
            "Chỉ hỗ trợ .docx / .pptx / .md — file “%s” chưa làm được."
            % command.filename
        )
    root = _doc_root_for(request)
    target = root / command.filename
    if command.kind == "edit" and not target.exists():
        return _message(
            "Chưa tìm thấy “%s” trong thư mục báo cáo. Kiểm tra lại tên file, "
            "hoặc dùng “tạo báo cáo %s” để tạo mới." % (command.filename, command.filename)
        )
    result = run_report_edit(
        command.filename,
        command.operations,
        create=(command.kind == "create"),
        command=request.question,
        doc_root=root,
    )
    if not result.get("ok"):
        return _message(
            "Chưa làm được: %s" % (result.get("error_vi") or "lỗi không rõ.")
        )
    summary = "Đã %s **%s**" % (
        "tạo" if command.kind == "create" else "sửa",
        Path(result["path"]).name,
    )
    return ChatActionOutcome(
        action=ACTION_NAME,
        title=ACTION_TITLE,
        blocks=(ChatActionBlock(BLOCK_MARKDOWN, text=_artifact_card(result, summary)),),
    )


def register() -> ChatAction:
    return register_action(
        ChatAction(
            name=ACTION_NAME,
            title=ACTION_TITLE,
            hints=_HINTS,
            handler=_handler,
            description=(
                "Tạo mới / sửa báo cáo .docx / .pptx / .md bằng lời trong chat; "
                "sửa file có sẵn luôn backup trước khi ghi; trả file đính kèm "
                "trong câu trả lời (UX-AGENT-REPORT)."
            ),
        )
    )


register()
